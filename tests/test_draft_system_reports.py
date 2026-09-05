from __future__ import annotations

import copy
import hashlib
import io
import json

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import update
from test_draft_constraint_review import base_case as _base_case
from test_draft_constraint_review import case as _case
from test_draft_constraint_review import inputs
from test_draft_scope import uid
from test_draft_scope_reports import _fast_renderers
from test_draft_system_matches import counts, create

from classifire.models import DraftScopeReport, User
from classifire.services.draft_scope import DraftScopeError
from classifire.services.draft_scope_reports import (
    _checksum,
    create_report,
    list_reports,
    read_report,
    report_bytes,
    report_freshness,
    validate_report_snapshot,
)
from classifire.services.draft_system_matches import (
    read_match_revision,
    save_constraint_review,
    save_review,
)

base_case = _base_case
case = _case


@pytest.mark.parametrize("measured", [False, True])
def test_saved_review_pair_restart_stale_and_no_downstream(case, measured, monkeypatch):
    with case["factory"]() as db:
        actor = db.get(User, uid(100))
        match = create(db, case)
        saved = read_match_revision(db, actor, case["draft_id"], match.id)
        decisions = copy.deepcopy(saved["decisions"])
        decisions[0].update(decision="keep", notes='=HYPERLINK("bad") <script> literal review')
        saved = save_review(db, actor, case["draft_id"], match.id, 1, decisions)
        if measured:
            saved = save_constraint_review(
                db,
                actor,
                case["draft_id"],
                match.id,
                saved["revision"],
                case["variant_id"],
                inputs(),
                storage_root=case["storage_root"],
            )
        before = counts(db)

        def forbidden(*args, **kwargs):
            pytest.fail("Reporting must not run retrieval, calculation or canonical writes")

        for target in (
            "classifire.services.draft_system_matches.create_match",
            "classifire.services.technical.search_variants",
            "classifire.services.calculation.recalculate_estimate",
            "classifire.services.snapshot.lock_snapshot",
        ):
            monkeypatch.setattr(target, forbidden)
        report = create_report(
            db,
            actor,
            case["draft_id"],
            2,
            match_id=match.id,
            match_revision=saved["revision"],
        )
        snapshot = read_report(db, actor, case["draft_id"], report.id)
        assert snapshot["system_match"] == saved
        assert snapshot["scope"] == saved["scope"]
        assert snapshot["profile"] == "scope-and-system"
        assert snapshot["render_version"] == 2
        assert not report_freshness(
            db,
            actor,
            case["draft_id"],
            report.id,
            storage_root=case["storage_root"],
        )
        pdf = report_bytes(db, actor, case["draft_id"], report.id, "pdf")
        xlsx = report_bytes(db, actor, case["draft_id"], report.id, "xlsx")
        text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf)).pages)
        assert "not technical approval" in text
        assert "Kept for review" in text
        assert "must-never-export-source-json" not in text
        workbook = load_workbook(io.BytesIO(xlsx))
        assert {"Review and coverage", "System candidates", "Measured limits"} <= set(
            workbook.sheetnames
        )
        cells = [
            cell for sheet in workbook for row in sheet for cell in row if cell.value is not None
        ]
        assert all(cell.data_type != "f" for cell in cells)
        for name in ("Review and coverage", "System candidates", "Measured limits"):
            assert workbook[name].row_dimensions[1].height >= 30
        values = "\n".join(str(cell.value) for cell in cells)
        assert decisions[0]["notes"] in values
        assert saved["sha256"] in values
        assert saved["candidates"][0]["fields_sha256"] in values
        assert "must-never-export-source-json" not in values
        assert ("within_limits" in values) is measured
        workbook.close()
        after = counts(db)
        assert {key: value for key, value in before.items() if key != "audit_events"} == {
            key: value for key, value in after.items() if key != "audit_events"
        }
        report_id, match_id = report.id, match.id
        db.commit()
    with case["factory"]() as db:
        actor = db.get(User, uid(100))
        save_review(db, actor, case["draft_id"], match_id, saved["revision"], decisions)
        db.commit()
        assert report_freshness(
            db, actor, case["draft_id"], report_id, storage_root=case["storage_root"]
        )
        assert read_report(db, actor, case["draft_id"], report_id) == snapshot
        assert report_bytes(db, actor, case["draft_id"], report_id, "pdf") == pdf
        assert report_bytes(db, actor, case["draft_id"], report_id, "xlsx") == xlsx


def test_permissions_source_binding_and_old_scope_profile(case, monkeypatch):
    _fast_renderers(monkeypatch)
    with case["factory"]() as db:
        actor = db.get(User, uid(100))
        match = create(db, case)
        old = create_report(db, actor, case["draft_id"], 2)
        report = create_report(db, actor, case["draft_id"], 2, match_id=match.id, match_revision=1)
        db.commit()
        with pytest.raises(DraftScopeError):
            create_report(db, actor, case["draft_id"], 1, match_id=match.id, match_revision=1)
        with pytest.raises(DraftScopeError):
            create_report(db, actor, case["draft_id"], 2, match_id=match.id)
        with pytest.raises(DraftScopeError):
            read_report(db, db.get(User, uid(101)), case["draft_id"], report.id)
        actor.role = "project_manager"
        db.commit()
        assert [row.id for row in list_reports(db, actor, case["draft_id"])] == [old.id]
        assert read_report(db, actor, case["draft_id"], old.id)["render_version"] == 1
        for action in (
            lambda: read_report(db, actor, case["draft_id"], report.id),
            lambda: report_bytes(db, actor, case["draft_id"], report.id, "pdf"),
        ):
            with pytest.raises(DraftScopeError) as denied:
                action()
            assert denied.value.status_code == 403


def test_source_change_warns_without_rewriting_and_tampered_snapshot_fails(case, monkeypatch):
    _fast_renderers(monkeypatch)
    with case["factory"]() as db:
        actor = db.get(User, uid(100))
        match = create(db, case)
        report = create_report(db, actor, case["draft_id"], 2, match_id=match.id, match_revision=1)
        snapshot = read_report(db, actor, case["draft_id"], report.id)
        pdf = report.pdf_bytes
        db.commit()
        case["source_path"].write_bytes(b"Changed synthetic source")
        assert report_freshness(
            db, actor, case["draft_id"], report.id, storage_root=case["storage_root"]
        )
        assert report_bytes(db, actor, case["draft_id"], report.id, "pdf") == pdf
        changed = copy.deepcopy(snapshot)
        changed["system_match"]["scope"]["revision"] += 1
        changed["sha256"] = _checksum(changed)
        with pytest.raises(DraftScopeError):
            validate_report_snapshot(changed)
        raw = json.dumps([])
        db.execute(
            update(DraftScopeReport)
            .where(DraftScopeReport.id == report.id)
            .values(
                snapshot_json=raw,
                snapshot_hash=hashlib.sha256(raw.encode()).hexdigest(),
            )
        )
        db.commit()
        with pytest.raises(DraftScopeError) as failure:
            read_report(db, actor, case["draft_id"], report.id)
        assert failure.value.code == "DRAFT_REPORT_INTEGRITY_FAILED"
