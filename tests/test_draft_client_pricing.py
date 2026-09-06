from __future__ import annotations

import hashlib
import io
import json
import re
from contextlib import asynccontextmanager
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import func, select
from test_draft_client import client_case as _client_case
from test_draft_client import confirm, tool
from test_draft_client_capabilities import propose, read, saved
from test_draft_estimates import add_payload
from test_draft_pdf_ui import pdf_app as _pdf_app
from test_draft_pdf_ui import pdf_setup as _pdf_setup
from test_draft_pricing import MAPPING, workbook_bytes
from test_draft_scope import sample_payload, uid
from test_draft_scope_ui import _app, _assert_no_canonical_scope, _csrf, _login
from test_draft_scope_ui import scope_password_hash as _password_hash
from test_shared_file_containment import postgresql_session_factory as _postgres_fixture

from classifire.draft_client import configure
from classifire.draft_client_auth import ESTIMATE, EXPORT, READ, WRITE
from classifire.models import (
    DraftClientRequest,
    DraftEstimate,
    DraftPricingSource,
    StoredFile,
    User,
)
from classifire.security import ROLE_PERMISSIONS
from classifire.services import draft_estimates as estimates
from classifire.services import draft_pricing_intake as pricing
from classifire.services import draft_scope as scopes
from classifire.services import draft_scope_reports as scope_reports

pdf_app = _pdf_app
pdf_setup = _pdf_setup
scope_password_hash = _password_hash
postgresql_session_factory = _postgres_fixture
client_case = _client_case


@pytest.fixture
def scope_app(pdf_app):
    """Reuse client identity setup with the real PostgreSQL containment fixture."""
    return SimpleNamespace(
        app=pdf_app.app,
        factory=pdf_app.factory,
        users={"owner": pdf_app.ids[0], "other": pdf_app.ids[1]},
    )


def retain_workbook(c, content, filename="synthetic.xlsx"):
    with c.scope.factory() as db:
        actor = db.get(User, c.scope.users["owner"])
        source = pricing.intake().retain(
            db, actor, c.draft_id, filename, content, settings=c.settings
        )
        db.commit()
        pricing.intake().scan_source(db, actor, c.draft_id, source.id, settings=c.settings)
        db.commit()
        source, rows = pricing.preview(
            db, actor, c.draft_id, source.id, 1, 1, MAPPING, settings=c.settings
        )
        stored = db.get(StoredFile, source.stored_file_id)
        return SimpleNamespace(
            id=source.id,
            document_hash=source.document_sha256,
            source_hash=source.source_sha256,
            scan_hash=hashlib.sha256(source.scan_json.encode()).hexdigest(),
            rows=rows,
            stored_id=source.stored_file_id,
            path=Path(stored.storage_path),
        )


@pytest.fixture
def pricing_case(client_case, pdf_app, monkeypatch):
    c = client_case
    c.settings = pdf_app.settings
    c.draft_id = pdf_app.ids[2]
    c.policy["clients"]["synthetic-client"] = [READ, WRITE, EXPORT, ESTIMATE]
    c.path.write_text(json.dumps(c.policy), encoding="utf-8")
    c.full_token = lambda: c.token(scope=" ".join([READ, WRITE, EXPORT, ESTIMATE]))
    for module in (
        "services.draft_client_capabilities",
        "draft_client_capability_tools",
        "draft_estimate_ui",
        "draft_pricing_ui",
        "draft_scope_ui",
    ):
        monkeypatch.setattr("classifire." + module + ".get_settings", lambda: c.settings)
    with c.scope.factory() as db:
        actor = db.get(User, c.scope.users["owner"])
        scopes.save_revision(db, actor, c.draft_id, 1, sample_payload())
        estimate = estimates.create_estimate(db, actor, c.draft_id, 2)
        content = estimates.add_line(db, actor, c.draft_id, estimate.id, 1, add_payload())
        c.estimate_id = estimate.id
        c.line_id = content["lines"][0]["line_id"]
        db.commit()
    c.source = retain_workbook(c, workbook_bytes())
    return c


def preview(client, c, *, source=None, **overrides):
    source = source or c.source
    return tool(
        client,
        c.full_token(),
        "preview_pricing_rows",
        {"draft_id": c.draft_id, "source_id": source.id, "mapping": MAPPING, **overrides},
    )


def operation(c, *, source=None, row=None, revision=2, **overrides):
    source = source or c.source
    row = row or source.rows[0]
    return {
        "action": "apply_workbook_rate",
        "draft_id": c.draft_id,
        "estimate_id": c.estimate_id,
        "expected_revision": revision,
        "line_id": c.line_id,
        "source_id": source.id,
        "sheet_index": 1,
        "header_row": 1,
        "mapping": MAPPING,
        "row_number": row["row"],
        "expected_document_hash": source.document_hash,
        "expected_row_hash": row["sha256"],
        "recovery_note": "Service work only; shared closure excluded",
        **overrides,
    }


def pending_rate(client, c, **kwargs):
    return tool(client, c.full_token(), "propose_capability", {"operation": operation(c, **kwargs)})


def unchanged(c, pending, revision=2):
    with c.scope.factory() as db:
        assert db.get(DraftEstimate, c.estimate_id).latest_revision == revision
        assert db.get(DraftClientRequest, pending["request_id"]).status == "pending"


def test_client_workbook_rate_keeps_exact_cells_history_reports_and_reopened_state(pricing_case):
    c = pricing_case
    with TestClient(c.scope.app, base_url="https://testserver") as client:
        _login(client)
        sources = tool(client, c.full_token(), "list_pricing_sources", {"draft_id": c.draft_id})
        assert sources["limit"] == 20
        assert sources["sources"][0]["id"] == c.source.id
        rows = preview(client, c)
        assert rows["mode"] == "mapped" and rows["rows"] == c.source.rows
        assert rows["source"]["document_sha256"] == c.source.document_hash
        assert rows["source"]["source_sha256"] == c.source.source_hash
        assert rows["source"]["scan_sha256"] == c.source.scan_hash
        assert rows["rows"][0]["fields"]["rate"]["address"] == "D2"
        original = read(client, c, "estimate", c.estimate_id)["artifact"]

        rejected = pending_rate(client, c)
        unchanged(c, rejected)
        assert read(client, c, "estimate", c.estimate_id)["artifact"] == original
        confirm(client, rejected, decision="reject")
        assert read(client, c, "estimate", c.estimate_id)["artifact"] == original
        pending = pending_rate(client, c)
        page = client.get(pending["review_url"])
        assert page.status_code == 200 and "D2" in page.text and "120.25" in page.text
        assert not re.search(r'<form[^>]+action="[^"]*/apply"', page.text)
        digest = re.search(r'name="payload_hash" value="([^"]+)"', page.text).group(1)
        unchanged(c, pending)
        saved(client, c, pending)
        assert (
            client.post(
                pending["review_url"],
                data={
                    "csrf_token": _csrf(page.text),
                    "payload_hash": digest,
                    "decision": "confirm",
                },
            ).status_code
            == 409
        )
        selected = read(client, c, "estimate", c.estimate_id)["artifact"]
        assert selected["revision"] == 3
        assert selected["lines"][0]["unit_sell_rate"] == "120.25"
        assert selected["lines"][0]["original_rate"] == "1.005"
        assert selected["summary"]["priced_subtotal_ex_tax"] == "240.50"
        assert len(selected["lines"][0]["history"]) == 2
        assert selected["pricing_sources"][0]["row"] == c.source.rows[0]
        assert selected["pricing_sources"][0]["approval_status"] == "unreviewed"
        detail = f"/scopes/{c.draft_id}/estimates/{c.estimate_id}"
        assert client.get(detail + "/download").json() == selected

        saved(
            client,
            c,
            propose(
                client,
                c,
                "override_estimate_line",
                estimate_id=c.estimate_id,
                expected_revision=3,
                line_id=c.line_id,
                override={
                    "quantity": "2",
                    "unit_sell_rate": "135",
                    "reason": "Reviewed manual adjustment",
                },
            ),
        )
        overridden = read(client, c, "estimate", c.estimate_id)["artifact"]
        assert overridden["lines"][0]["original_rate"] == "1.005"
        assert overridden["lines"][0]["unit_sell_rate"] == "135"
        assert overridden["pricing_sources"] == selected["pricing_sources"]
        assert len(overridden["lines"][0]["history"]) == 3
        downloads = []
        for profile in ("estimate-only", "complete"):
            report = saved(
                client,
                c,
                propose(
                    client,
                    c,
                    "estimate_report",
                    estimate_id=c.estimate_id,
                    estimate_revision=4,
                    profile=profile,
                ),
            )["report_id"]
            assert (
                read(client, c, "estimate-report", report, estimate_id=c.estimate_id)["artifact"][
                    "estimate"
                ]
                == overridden
            )
            for fmt in ("pdf", "xlsx"):
                info = tool(
                    client,
                    c.full_token(),
                    "report_download",
                    {
                        "draft_id": c.draft_id,
                        "estimate_id": c.estimate_id,
                        "report_id": report,
                        "format_name": fmt,
                    },
                )
                response = client.get(
                    info["authenticated_url"], headers={"Authorization": "Bearer " + c.full_token()}
                )
                assert response.status_code == 200
                assert hashlib.sha256(response.content).hexdigest() == info["sha256"]
                assert client.get(info["browser_url"]).content == response.content
                if fmt == "pdf":
                    assert "D2" in "\n".join(
                        page.extract_text()
                        for page in PdfReader(io.BytesIO(response.content)).pages
                    )
                else:
                    book = load_workbook(io.BytesIO(response.content), data_only=False)
                    cells = [cell for row in book["Pricing sources"] for cell in row]
                    assert "D2" in [str(cell.value) for cell in cells]
                    assert "120.25 (number)" in [str(cell.value) for cell in cells]
                    assert not any(cell.data_type == "f" for cell in cells)
                    book.close()
                downloads.append((info, response.content))
        assert read(client, c, "estimate", c.estimate_id, revision=2)["artifact"] == original
        assert read(client, c, "estimate", c.estimate_id, revision=3)["artifact"] == selected

    app = _app(c.scope.factory)
    mounted = configure(app, c.authority, c.scope.factory)

    @asynccontextmanager
    async def lifespan(app):
        async with mounted.router.lifespan_context(mounted):
            yield

    app.router.lifespan_context = lifespan
    with TestClient(app, base_url="https://testserver") as restarted:
        assert read(restarted, c, "estimate", c.estimate_id)["artifact"] == overridden
        for info, content in downloads:
            assert (
                restarted.get(
                    info["authenticated_url"], headers={"Authorization": "Bearer " + c.full_token()}
                ).content
                == content
            )
    _assert_no_canonical_scope(c.scope.factory)


def test_pricing_preview_pages_verified_cells_without_truncating_mapped_provenance(pricing_case):
    c = pricing_case
    book = load_workbook(io.BytesIO(workbook_bytes()))
    sheet = book["Rates"]
    sheet["B2"] = "Synthetic long description " + "x" * 300
    for row_number in range(5, 10):
        sheet.append([f"SYN-{row_number}", "Synthetic page", "each", 12, "AUD", "excluded"])
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    source = retain_workbook(c, stream.getvalue(), "paged.xlsx")
    with TestClient(c.scope.app, base_url="https://testserver") as client:
        first = preview(client, c, source=source, mapping=None)
        assert first["mode"] == "unmapped" and first["limit"] == 5
        assert first["sheets"] == [{"index": 1, "name": "Rates", "rows": 9, "columns": 11}]
        assert [row["row"] for row in first["rows"]] == [1, 2, 3, 4, 5]
        long_cell = next(cell for cell in first["rows"][1]["cells"] if cell["address"] == "B2")
        assert len(long_cell["value"]) == 200 and long_cell["truncated"] is True
        assert long_cell["original_length"] == len(sheet["B2"].value)
        second = preview(client, c, source=source, mapping=None, after_row=first["next_after_row"])
        assert [row["row"] for row in second["rows"]] == [6, 7, 8, 9]
        assert second["next_after_row"] is None
        mapped = preview(client, c, source=source)
        assert mapped["rows"] == source.rows[:5]
        assert len(mapped["rows"][0]["fields"]["description"]["value"]) > 200
        assert mapped["rows"][2]["fields"]["rate"]["kind"] == "formula"
        assert mapped["rows"][2]["problems"]
        rest = preview(client, c, source=source, after_row=mapped["next_after_row"])
        assert mapped["rows"] + rest["rows"] == source.rows
        assert rest["next_after_row"] is None
        for invalid in (True, "1", -1, 10):
            tool(
                client,
                c.full_token(),
                "preview_pricing_rows",
                {"draft_id": c.draft_id, "source_id": source.id, "after_row": invalid},
                error=True,
            )
        with c.scope.factory() as db:
            assert db.scalar(select(func.count()).select_from(DraftClientRequest)) == 0
            assert db.get(DraftEstimate, c.estimate_id).latest_revision == 2


def test_pricing_client_permissions_owner_and_history_are_current(pricing_case, monkeypatch):
    c = pricing_case
    with TestClient(c.scope.app, base_url="https://testserver") as client:
        _login(client)
        for name, args in (
            ("list_pricing_sources", {"draft_id": c.draft_id}),
            ("preview_pricing_rows", {"draft_id": c.draft_id, "source_id": c.source.id}),
            ("propose_capability", {"operation": operation(c)}),
        ):
            assert "CLIENT_AUTHORIZATION_REQUIRED" in str(
                tool(client, c.token(), name, args, error=True)
            )
            foreign = c.token(user="other", scope=" ".join([READ, WRITE, EXPORT, ESTIMATE]))
            assert "DRAFT_NOT_FOUND" in str(tool(client, foreign, name, args, error=True))
        pending = pending_rate(client, c)
        saved(client, c, pending)
        current = pending_rate(client, c, revision=3)
        with monkeypatch.context() as rights:
            rights.setitem(
                ROLE_PERMISSIONS, "estimator", ROLE_PERMISSIONS["estimator"] - {"library:read"}
            )
            assert client.get(current["review_url"]).status_code == 403
            assert client.get(pending["review_url"]).status_code == 403
            for name, args in (
                ("list_pricing_sources", {"draft_id": c.draft_id}),
                ("preview_pricing_rows", {"draft_id": c.draft_id, "source_id": c.source.id}),
                (
                    "read_capability_artifact",
                    {"draft_id": c.draft_id, "kind": "estimate", "artifact_id": c.estimate_id},
                ),
            ):
                tool(client, c.full_token(), name, args, error=True)
        c.policy["clients"]["synthetic-client"].remove(ESTIMATE)
        c.path.write_text(json.dumps(c.policy), encoding="utf-8")
        assert client.get(current["review_url"]).status_code == 403
        c.policy["clients"]["synthetic-client"].append(ESTIMATE)
        c.policy["revoked_token_ids"] = ["synthetic-token"]
        c.path.write_text(json.dumps(c.policy), encoding="utf-8")
        assert client.get(current["review_url"]).status_code == 403
        unchanged(c, current, 3)


def test_pricing_proposal_rejects_changed_hash_mapping_target_and_stale_estimate(pricing_case):
    c = pricing_case
    with TestClient(c.scope.app, base_url="https://testserver") as client:
        _login(client)
        for changed in (
            {"expected_document_hash": "0" * 64},
            {"expected_row_hash": "0" * 64},
            {"mapping": {**MAPPING, "rate": 1}},
            {"mapping": {**MAPPING, "rate": "4"}},
            {"mapping": {**MAPPING, "rate": True}},
            {"mapping": {**MAPPING, "approval": 12}},
            {"mapping": {key: value for key, value in MAPPING.items() if key != "rate"}},
            {"expected_document_hash": "A" * 64},
            {"line_id": uid(999)},
        ):
            tool(
                client,
                c.full_token(),
                "propose_capability",
                {"operation": operation(c, **changed)},
                error=True,
            )
        with c.scope.factory() as db:
            assert db.scalar(select(func.count()).select_from(DraftClientRequest)) == 0
        pending = pending_rate(client, c)
        saved(
            client,
            c,
            propose(
                client,
                c,
                "override_estimate_line",
                estimate_id=c.estimate_id,
                expected_revision=2,
                line_id=c.line_id,
                override={"quantity": "2", "unit_sell_rate": "3", "reason": "Newer human edit"},
            ),
        )
        assert client.get(pending["review_url"]).status_code == 409
        unchanged(c, pending, 3)


@pytest.mark.parametrize("change", ["source_bytes", "scan", "document", "quarantine"])
def test_pending_workbook_rate_binds_current_retained_source_scan_and_document(
    pricing_case, change
):
    c = pricing_case
    with TestClient(c.scope.app, base_url="https://testserver") as client:
        _login(client)
        pending = pending_rate(client, c)
        page = client.get(pending["review_url"])
        digest = re.search(r'name="payload_hash" value="([^"]+)"', page.text).group(1)
        with c.scope.factory() as db:
            source = db.get(DraftPricingSource, c.source.id)
            if change == "source_bytes":
                c.source.path.write_bytes(b"changed synthetic workbook bytes")
            elif change == "scan":
                pricing.intake().scan_source(
                    db,
                    db.get(User, c.scope.users["owner"]),
                    c.draft_id,
                    source.id,
                    settings=c.settings,
                )
                assert hashlib.sha256(source.scan_json.encode()).hexdigest() != c.source.scan_hash
            elif change == "document":
                source.document_json = "{}"
            else:
                db.get(StoredFile, c.source.stored_id).malware_scan_status = "malware_detected"
            db.commit()
        response = client.post(
            pending["review_url"],
            data={"csrf_token": _csrf(page.text), "payload_hash": digest, "decision": "confirm"},
        )
        assert response.status_code == 409, response.text
        assert (
            "CLIENT_INPUT_CHANGED" if change == "scan" else "PRICING_XLSX_SOURCE_NOT_READY"
        ) in response.text
        unchanged(c, pending)


def test_formula_and_incompatible_units_leave_pricing_request_pending(pricing_case):
    c = pricing_case
    mismatch = retain_workbook(c, workbook_bytes(unit="m"), "metre-rate.xlsx")
    unsupported = retain_workbook(c, workbook_bytes(unit="box"), "box-rate.xlsx")
    with TestClient(c.scope.app, base_url="https://testserver") as client:
        _login(client)
        for source, row, expected in (
            (c.source, c.source.rows[2], "PRICING_RATE_UNRESOLVED"),
            (mismatch, mismatch.rows[0], "PRICING_UNIT_MISMATCH"),
            (unsupported, unsupported.rows[0], "PRICING_RATE_UNRESOLVED"),
        ):
            pending = pending_rate(client, c, source=source, row=row)
            assert expected in confirm(client, pending, status=422).text
            unchanged(c, pending)


def test_existing_manual_pending_input_binding_remains_compatible(pricing_case):
    c = pricing_case
    with TestClient(c.scope.app, base_url="https://testserver") as client:
        _login(client)
        original = read(client, c, "estimate", c.estimate_id)["artifact"]
        pending = propose(
            client,
            c,
            "override_estimate_line",
            estimate_id=c.estimate_id,
            expected_revision=2,
            line_id=c.line_id,
            override={
                "quantity": "2",
                "unit_sell_rate": "3",
                "reason": "Existing pending manual correction",
            },
        )
        with c.scope.factory() as db:
            actor = db.get(User, c.scope.users["owner"])
            old_inputs = {
                "scope": None,
                "match": None,
                "estimate": original,
                "release": None,
                "project": scope_reports._project(db, scopes.get_draft(db, actor, c.draft_id)),
            }
            expected_hash = hashlib.sha256(
                json.dumps(
                    old_inputs, sort_keys=True, separators=(",", ":"), ensure_ascii=True
                ).encode()
            ).hexdigest()
            stored = json.loads(db.get(DraftClientRequest, pending["request_id"]).payload_json)
            assert stored["input_hash"] == expected_hash
        saved(client, c, pending)
        changed = read(client, c, "estimate", c.estimate_id)["artifact"]
        assert changed["lines"][0]["unit_sell_rate"] == "3"
        assert changed.get("pricing_sources", []) == []


def test_manual_estimate_workbook_request_history_requires_current_library_rights(
    pricing_case, monkeypatch
):
    c = pricing_case
    with TestClient(c.scope.app, base_url="https://testserver") as client:
        _login(client)
        original = read(client, c, "estimate", c.estimate_id)["artifact"]
        assert original["revision"] == 2 and original.get("pricing_sources", []) == []
        pending = pending_rate(client, c)
        rejected = pending_rate(client, c)
        confirm(client, rejected, decision="reject")
        assert client.get(rejected["review_url"]).status_code == 200

        with monkeypatch.context() as rights:
            rights.setitem(
                ROLE_PERMISSIONS, "estimator", ROLE_PERMISSIONS["estimator"] - {"library:read"}
            )
            # Manual Estimate access still works; its reader cannot guard workbook
            # proposal details because no workbook rate was ever applied.
            assert read(client, c, "estimate", c.estimate_id)["artifact"] == original
            for request in (pending, rejected):
                assert client.get(request["review_url"]).status_code == 403

        unchanged(c, pending)
        assert read(client, c, "estimate", c.estimate_id)["artifact"] == original
        with c.scope.factory() as db:
            assert db.get(DraftClientRequest, rejected["request_id"]).status == "rejected"
