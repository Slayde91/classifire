from __future__ import annotations

import hashlib
import io
import json
from datetime import UTC, datetime
from pathlib import Path
from types import SimpleNamespace

import pytest
from reportlab.pdfgen.canvas import Canvas
from sqlalchemy import func, select, update
from test_shared_file_containment import postgresql_session_factory as _postgres_fixture

from classifire.config import Settings
from classifire.models import DraftPdfSource, Estimate, Opening, Service, StoredFile, User
from classifire.services import draft_pdf_intake as intake
from classifire.services import draft_scope as scope
from classifire.services.draft_scope_evidence import reference_status
from classifire.services.malware_scan import MalwareScanError, ScanVerdict

postgresql_session_factory = _postgres_fixture


def pdf_bytes(label="Synthetic evidence only"):
    output = io.BytesIO()
    canvas = Canvas(output)
    canvas.drawString(60, 760, label)
    canvas.drawString(60, 730, "Shared opening: service count and dimensions remain unresolved.")
    canvas.save()
    return output.getvalue()


@pytest.fixture
def pdf_setup(postgresql_session_factory, tmp_path, monkeypatch):
    factory = postgresql_session_factory
    settings = Settings(storage_root=tmp_path / "retained", clamav_host="127.0.0.1", _env_file=None)
    settings.storage_root.mkdir()
    with factory() as db:
        owner = User(
            email="pdf-owner@example.test",
            full_name="PDF owner",
            role="estimator",
            is_active=True,
            password_hash="unused",  # noqa: S106 - synthetic non-authenticating fixture
        )  # noqa: S106
        foreign = User(
            email="pdf-other@example.test",
            full_name="Other owner",
            role="estimator",
            is_active=True,
            password_hash="unused",  # noqa: S106 - synthetic non-authenticating fixture
        )  # noqa: S106
        db.add_all([owner, foreign])
        db.flush()
        draft = scope.create_draft_project(db, owner, "PDF-SYNTHETIC", "PDF review prototype")
        db.commit()
        ids = (owner.id, foreign.id, draft.id)

    def clean(content, **kwargs):
        now = datetime.now(UTC).isoformat()
        return ScanVerdict(
            hashlib.sha256(content).hexdigest(),
            len(content),
            "clean",
            "ClamAV 1.5.4",
            "28108",
            now,
            now,
        )

    monkeypatch.setattr(intake.malware_scan, "scan_bytes", clean)
    return SimpleNamespace(factory=factory, settings=settings, ids=ids, clean=clean)


def retained(db, setup):
    actor = db.get(User, setup.ids[0])
    source = intake.retain_pdf(
        db, actor, setup.ids[2], "synthetic.pdf", pdf_bytes(), settings=setup.settings
    )
    db.commit()
    return actor, source


def prepared(db, setup):
    actor, source = retained(db, setup)
    intake.scan_source(db, actor, setup.ids[2], source.id, settings=setup.settings)
    db.commit()
    return actor, source


def test_retained_pdf_review_and_history_survive_new_session_and_import(pdf_setup):
    x = pdf_setup
    with x.factory() as db:
        actor, source = retained(db, x)
        old = scope.revision_bytes(db, actor, x.ids[2])
        assert intake.source_info(db, actor, x.ids[2], source.id)["status"] == "pending"
        with pytest.raises(scope.DraftScopeError):
            intake.page_preview(db, actor, x.ids[2], source.id, 1, settings=x.settings)
        intake.scan_source(db, actor, x.ids[2], source.id, settings=x.settings)
        db.commit()
        doc = intake.read_document(db, actor, x.ids[2], source.id, settings=x.settings)
        assert "Synthetic evidence" in doc["pages"][0]["text"]
        assert intake.page_preview(
            db, actor, x.ids[2], source.id, 1, settings=x.settings
        ).startswith(b"\x89PNG")
        saved = intake.review_page(
            db,
            actor,
            x.ids[2],
            source.id,
            1,
            1,
            "Shared opening exists; service count is unresolved.",
            "Unresolved",
            source.document_sha256,
            settings=x.settings,
        )
        assert saved["schema_version"] == "CLASSIFIRE-DRAFT-SCOPE-v3"
        assert len(saved["evidence_refs"]) == 1
        assert scope.validate_portable_artifact(scope._json(saved)) == saved
        original = scope.revision_bytes(db, actor, x.ids[2], 2)
        edited = json.loads(json.dumps(saved["content"]))
        edited["observations"][0]["text"] = "Changed observation requires another review"
        later = scope.save_revision(db, actor, x.ids[2], 2, edited)
        assert "changed" in reference_status(
            later["evidence_refs"][0], later["content"]["observations"]
        )
        assert scope.revision_bytes(db, actor, x.ids[2], 1) == old
        target = scope.create_draft_project(db, actor, "IMPORT-PDF", "Unverified source import")
        imported = scope.apply_import(
            db,
            actor,
            target.id,
            1,
            original,
            expected_source_hash=hashlib.sha256(original).hexdigest(),
        )
        assert imported["evidence_refs"][0]["origin"] == "imported_unverified"
        assert "SCOPE_SOURCE_UNVERIFIED" in intake.scope_evidence_staleness(
            db, actor, target.id, imported, storage_root=x.settings.storage_root
        )
        db.commit()
        source_id = source.id
    with x.factory() as db:
        actor = db.get(User, x.ids[0])
        assert scope.revision_bytes(db, actor, x.ids[2], 2) == original
        assert intake.read_document(db, actor, x.ids[2], source_id, settings=x.settings) == doc
        for model in (Estimate, Opening, Service):
            assert db.scalar(select(func.count()).select_from(model)) == 0


@pytest.mark.parametrize("status", ["error", "malware"])
def test_scan_failure_and_malware_remain_blocked(pdf_setup, monkeypatch, status):
    x = pdf_setup
    with x.factory() as db:
        actor, source = retained(db, x)

        def verdict(content, **kwargs):
            if status == "error":
                raise MalwareScanError("SCAN_TIMEOUT")
            result = x.clean(content)
            return ScanVerdict(
                result.sha256,
                result.size_bytes,
                "malware_detected",
                result.engine,
                result.database_version,
                result.database_date,
                result.scanned_at,
            )

        monkeypatch.setattr(intake.malware_scan, "scan_bytes", verdict)
        intake.scan_source(db, actor, x.ids[2], source.id, settings=x.settings)
        db.commit()
        assert intake.source_info(db, actor, x.ids[2], source.id)["status"] == (
            "scan_error" if status == "error" else "malware_detected"
        )
        with pytest.raises(scope.DraftScopeError):
            intake.read_document(db, actor, x.ids[2], source.id, settings=x.settings)
        if status == "malware":
            with pytest.raises(scope.DraftScopeError, match="PDF_SOURCE_QUARANTINED"):
                intake.scan_source(db, actor, x.ids[2], source.id, settings=x.settings)


def test_owner_cs_source_integrity_and_stale_review_refuse(pdf_setup):
    x = pdf_setup
    with x.factory() as db:
        actor, source = prepared(db, x)
        foreign = db.get(User, x.ids[1])
        with pytest.raises(scope.DraftScopeError):
            intake.read_document(db, foreign, x.ids[2], source.id, settings=x.settings)
        with pytest.raises(scope.DraftScopeError, match="PDF_REVIEW_SOURCE_CHANGED"):
            intake.review_page(
                db,
                actor,
                x.ids[2],
                source.id,
                1,
                1,
                "Unresolved",
                "Unresolved",
                "0" * 64,
                settings=x.settings,
            )
        stored = db.get(StoredFile, source.stored_file_id)
        Path(stored.storage_path).write_bytes(b"changed synthetic bytes")
        with pytest.raises(scope.DraftScopeError):
            intake.page_preview(db, actor, x.ids[2], source.id, 1, settings=x.settings)
        assert scope.read_revision(db, actor, x.ids[2])["revision"] == 1


def test_revocation_during_scan_never_releases_clean_content(pdf_setup, monkeypatch):
    x = pdf_setup
    with x.factory() as db:
        actor, source = retained(db, x)

        def revoke(content, **kwargs):
            db.execute(update(User).where(User.id == actor.id).values(is_active=False))
            db.flush()
            return x.clean(content)

        monkeypatch.setattr(intake.malware_scan, "scan_bytes", revoke)
        result = intake.scan_source(db, actor, x.ids[2], source.id, settings=x.settings)
        db.commit()
        assert result["status"] == "scan_error"
        assert db.get(StoredFile, source.stored_file_id).malware_scan_status == "scan_error"
        assert db.get(DraftPdfSource, source.id).document_json is None


@pytest.mark.parametrize("phase", ["malware_after_revocation", "parser_failure", "expired_scan"])
def test_failure_states_never_allow_page_review(pdf_setup, monkeypatch, phase):
    from dataclasses import replace
    from datetime import timedelta

    x = pdf_setup
    with x.factory() as db:
        actor, source = retained(db, x)
        if phase == "malware_after_revocation":

            def verdict(content, **kwargs):
                db.execute(update(User).where(User.id == actor.id).values(is_active=False))
                return replace(x.clean(content), status="malware_detected")

            monkeypatch.setattr(intake.malware_scan, "scan_bytes", verdict)
        elif phase == "parser_failure":

            def fail(*args):
                raise scope.DraftScopeError("PDF_PROCESSING_FAILED", 422)

            monkeypatch.setattr(intake, "_process", fail)
        intake.scan_source(db, actor, x.ids[2], source.id, settings=x.settings)
        db.commit()
        if phase == "malware_after_revocation":
            assert (
                db.get(StoredFile, source.stored_file_id).malware_scan_status == "malware_detected"
            )
        elif phase == "parser_failure":
            assert source.processing_error == "PDF_PROCESSING_FAILED"
            assert source.document_json is None
        else:
            value = json.loads(source.scan_json)
            value["database_date"] = (datetime.now(UTC) - timedelta(days=8)).isoformat()
            source.scan_json = json.dumps(value)
            db.commit()
        with pytest.raises(scope.DraftScopeError):
            intake.page_preview(db, actor, x.ids[2], source.id, 1, settings=x.settings)


def test_same_bytes_cannot_be_adopted_into_another_draft(pdf_setup):
    x = pdf_setup
    with x.factory() as db:
        actor = db.get(User, x.ids[0])
        content = pdf_bytes()
        first = intake.retain_pdf(db, actor, x.ids[2], "one.pdf", content, settings=x.settings)
        db.commit()
        assert (
            intake.retain_pdf(db, actor, x.ids[2], "renamed.pdf", content, settings=x.settings).id
            == first.id
        )
        other = scope.create_draft_project(db, actor, "SECOND", "Different source owner")
        db.commit()
        with pytest.raises(scope.DraftScopeError, match="PDF_UPLOAD_CONFLICT"):
            intake.retain_pdf(db, actor, other.id, "one.pdf", content, settings=x.settings)
        assert list((x.settings.storage_root / ".incoming").iterdir()) == []


def test_concurrent_quarantine_blocks_active_page_read(pdf_setup):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event

    from classifire.services.storage import quarantine_stored_file_bytes_for_update

    x = pdf_setup
    with x.factory() as db:
        actor, source = prepared(db, x)
        source_id, file_id = source.id, source.stored_file_id
        digest, size = source.source_sha256, source.source_size_bytes
    started = Event()

    def reader():
        with x.factory() as db:
            actor = db.get(User, x.ids[0])
            started.set()
            try:
                intake.read_document(db, actor, x.ids[2], source_id, settings=x.settings)
            except scope.DraftScopeError as exc:
                return exc.code
            return "unexpected read"

    with x.factory() as quarantine:
        quarantine.begin()
        quarantine_stored_file_bytes_for_update(
            quarantine, stored_file_id=file_id, observed_sha256=digest, observed_size_bytes=size
        )
        with ThreadPoolExecutor(max_workers=1) as pool:
            result = pool.submit(reader)
            assert started.wait(5)
            assert not result.done()
            quarantine.commit()
            assert result.result(timeout=10) == "PDF_SOURCE_NOT_READY"


def test_page_provenance_reaches_reports_and_quarantine_marks_downstream_stale(pdf_setup):
    import pymupdf
    from openpyxl import load_workbook

    from classifire.services import draft_estimate_reports as estimate_reports
    from classifire.services import draft_estimates as estimates
    from classifire.services import draft_scope_reports as reports
    from classifire.services.storage import quarantine_stored_file_bytes_for_update

    x = pdf_setup
    with x.factory() as db:
        actor, source = prepared(db, x)
        saved = intake.review_page(
            db,
            actor,
            x.ids[2],
            source.id,
            1,
            1,
            "Unresolved illustrated opening",
            "Unresolved",
            source.document_sha256,
            settings=x.settings,
        )
        report = reports.create_report(db, actor, x.ids[2], 2)
        estimate = estimates.create_estimate(db, actor, x.ids[2], 2)
        cost_report = estimate_reports.create_report(db, actor, x.ids[2], estimate.id, 1)
        originals = {
            kind: reports.report_bytes(db, actor, x.ids[2], report.id, kind)
            for kind in ("pdf", "xlsx")
        }
        with pymupdf.open(stream=originals["pdf"], filetype="pdf") as pdf:
            assert source.original_filename in "".join(page.get_text() for page in pdf)
        from zipfile import ZipFile

        from classifire.outputs.draft_branding import supplied_logo_path

        with ZipFile(io.BytesIO(originals["xlsx"])) as archive:
            assert archive.read("xl/media/image1.png") == supplied_logo_path().read_bytes()
        workbook = load_workbook(io.BytesIO(originals["xlsx"]))
        assert "Page Review References" in workbook.sheetnames
        assert source.source_sha256 in str(list(workbook["Page Review References"].values))
        estimate_pdf = estimate_reports.report_bytes(
            db, actor, x.ids[2], estimate.id, cost_report.id, "pdf"
        )
        with pymupdf.open(stream=estimate_pdf, filetype="pdf") as pdf:
            assert source.original_filename in "".join(page.get_text() for page in pdf)
        assert not reports.report_freshness(
            db, actor, x.ids[2], report.id, storage_root=x.settings.storage_root
        )
        assert not estimates.estimate_staleness(
            db, actor, x.ids[2], estimate.id, storage_root=x.settings.storage_root
        )
        quarantine_stored_file_bytes_for_update(
            db,
            stored_file_id=source.stored_file_id,
            observed_sha256=source.source_sha256,
            observed_size_bytes=source.source_size_bytes,
        )
        db.commit()
        assert reports.report_freshness(
            db, actor, x.ids[2], report.id, storage_root=x.settings.storage_root
        )
        assert "SCOPE_SOURCE_UNAVAILABLE" in estimates.estimate_staleness(
            db, actor, x.ids[2], estimate.id, storage_root=x.settings.storage_root
        )
        for kind, content in originals.items():
            assert reports.report_bytes(db, actor, x.ids[2], report.id, kind) == content
        assert scope.read_revision(db, actor, x.ids[2]) == saved
