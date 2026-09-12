from __future__ import annotations

import io
import json
from copy import deepcopy

import pytest
from openpyxl import load_workbook
from test_draft_constraint_review import base_case as _base_case
from test_draft_constraint_review import case as _case
from test_draft_constraint_review import inputs
from test_draft_entity_evidence_outputs import report_snapshot, review
from test_draft_scope import uid
from test_draft_scope_outputs import _id, _pdf_text, _seal, _snapshot
from test_draft_system_matches import create
from test_draft_word_evidence_outputs import word_scope

from classifire.models import User
from classifire.outputs import draft_scope as outputs
from classifire.services import draft_scope_reports as reports
from classifire.services import draft_system_match_contract as match_contract
from classifire.services import draft_system_matches as matches

base_case = _base_case
case = _case


def collection_snapshot(*, word=True):
    scope = word_scope() if word else _snapshot(imported=False)["scope"]
    selected = []
    for number, opening, service in ((32, 5, 8), (33, 6, 8), (34, 7, None)):
        saved = review(scope)
        saved.update(
            artifact_id=_id(number),
            target=match_contract.target_for(
                scope, _id(opening), _id(service) if service else None
            ),
        )
        saved["retrieval"]["missing_criteria"] = ["service_material", "frl"]
        selected.append(_seal(saved))
    snapshot = report_snapshot("scope-and-system", scope)
    snapshot.pop("system_match")
    snapshot.update(
        schema_version=reports.MULTI_SYSTEM_REPORT_SCHEMA_VERSION,
        render_version=11,
        system_matches=selected,
    )
    return _seal(snapshot)


def workbook_values(book):
    return [cell for sheet in book for row in sheet for cell in row if cell.value is not None]


@pytest.mark.parametrize("word", [False, True], ids=["manual-scope", "word-v7"])
def test_three_exact_rows_shared_service_blank_unknowns_and_word_export_parity(word, tmp_path):
    snapshot = collection_snapshot(word=word)
    reports.validate_report_snapshot(snapshot)
    original = deepcopy(snapshot)
    pdf = outputs.render_scope_report_pdf(snapshot)
    xlsx = outputs.render_scope_report_xlsx(snapshot)
    assert snapshot == original
    assert len(pdf) < reports.MAX_OUTPUT_BYTES and len(xlsx) < reports.MAX_OUTPUT_BYTES
    text = " ".join(_pdf_text(pdf).split())
    book = load_workbook(io.BytesIO(xlsx), data_only=False)
    cells = workbook_values(book)
    strings = " ".join(str(cell.value) for cell in cells)
    assert all(cell.data_type != "f" and cell.hyperlink is None for cell in cells)
    rows = {}
    for _, key, field, value in book["Selected reviews"].iter_rows(min_row=7, values_only=True):
        rows.setdefault(key, {})[field] = value
    assert set(rows) == {
        f"{match['artifact_id']} / revision {match['revision']}"
        for match in snapshot["system_matches"]
    }
    for match in snapshot["system_matches"]:
        fields = rows[f"{match['artifact_id']} / revision {match['revision']}"]
        assert fields["Opening ID"] == match["target"]["opening_id"]
        assert fields["Review ID"] == match["artifact_id"]
        assert fields["Review revision"] == str(match["revision"])
        assert fields["Review SHA-256"] == match["sha256"]
        assert fields["Scope SHA-256"] == snapshot["scope"]["sha256"]
        assert fields["Service ID"] == (match["target"]["service_id"] or "None - blank opening")
        for exact in (match["artifact_id"], match["sha256"], match["target"]["opening_id"]):
            assert exact in text and exact in strings
    assert rows[f"{_id(32)} / revision 1"]["Service ID"] == _id(8)
    assert rows[f"{_id(33)} / revision 1"]["Service ID"] == _id(8)
    assert rows[f"{_id(34)} / revision 1"]["Row kind"] == "Blank opening"
    for expected in (
        "Blank opening - no Service",
        "Unknown / unavailable",
        "No measured-limit review saved",
        "not technical approval",
        "review records are not quantities or counts of physical work",
        "No Estimate, pricing or release",
        'Do not run <img src="https://example.invalid/source.png"/>',
        snapshot["sha256"],
    ):
        assert expected in text and expected in strings
    # A shared Service is listed once, with separate physical links; no blank Service is invented.
    service_rows = list(book["Services"].iter_rows(min_row=7, values_only=True))
    assert sum(row[1] == _id(8) for row in service_rows) == 1
    assert next(row for row in service_rows if row[1] == _id(8))[4] == 2.125
    links = list(book["Service Links"].iter_rows(min_row=7, values_only=True))
    assert {(row[1], row[2]) for row in links if row[1] == _id(8)} == {
        (_id(8), _id(5)),
        (_id(8), _id(6)),
    }
    assert all(row[2] != _id(7) for row in links)
    if word:
        for expected in (
            "Synthetic Word.docx",
            "body-2/row-1/cell-1/p-1",
            "untrusted Word text",
            "word/media/image1.png",
            "synthetic-page-suggester",
            "Imported source and review claims are unverified",
        ):
            assert expected in text and expected in strings
    for name in ("Selected reviews", "Review and coverage", "System candidates", "Measured limits"):
        assert book[name].freeze_panes == "C7"
        assert book[name].auto_filter.ref
    book.close()
    (tmp_path / "three-review-report.pdf").write_bytes(pdf)
    (tmp_path / "three-review-report.xlsx").write_bytes(xlsx)
    (tmp_path / "three-review-snapshot.json").write_text(json.dumps(snapshot, ensure_ascii=False))


def test_saved_candidate_decisions_constraints_and_source_provenance_survive_pair(
    case, monkeypatch, tmp_path
):
    with case["factory"]() as db:
        actor = db.get(User, uid(100))
        refs = []
        originals = []
        for index, (opening, service) in enumerate(
            ((uid(2), uid(5)), (uid(3), uid(5)), (uid(4), None))
        ):
            row = create(db, case, opening_id=opening, service_id=service)
            saved = matches.read_match_revision(db, actor, case["draft_id"], row.id, 1)
            decisions = deepcopy(saved["decisions"])
            for decision in decisions:
                decision.update(
                    decision="keep", notes=f'ROW-{index} <b>literal</b> =HYPERLINK("bad")'
                )
            saved = matches.save_review(db, actor, case["draft_id"], row.id, 1, decisions)
            if index == 0:
                saved = matches.save_constraint_review(
                    db,
                    actor,
                    case["draft_id"],
                    row.id,
                    saved["revision"],
                    case["variant_id"],
                    inputs(None, None, None),
                    storage_root=case["storage_root"],
                )
            refs.append({"match_id": row.id, "match_revision": saved["revision"]})
            originals.append(saved)

        def forbidden(*args, **kwargs):
            pytest.fail("Rendering must not rerun matching, pricing, providers or canonical writes")

        for target in (
            "classifire.services.draft_system_matches.create_match",
            "classifire.services.technical.search_variants",
            "classifire.services.calculation.recalculate_estimate",
            "classifire.services.snapshot.lock_snapshot",
        ):
            monkeypatch.setattr(target, forbidden)
        report = reports.create_report(db, actor, case["draft_id"], 2, matches=list(reversed(refs)))
        saved_report = reports.read_report(db, actor, case["draft_id"], report.id)
        assert saved_report["system_matches"] == sorted(
            originals, key=lambda item: item["artifact_id"]
        )
        pdf, xlsx = report.pdf_bytes, report.xlsx_bytes
        text = " ".join(_pdf_text(pdf).split())
        book = load_workbook(io.BytesIO(xlsx), data_only=False)
        cells = workbook_values(book)
        strings = " ".join(str(cell.value) for cell in cells)
        for saved in originals:
            assert saved["sha256"] in text and saved["sha256"] in strings
            for candidate in saved["candidates"]:
                assert candidate["fields_sha256"] in text and candidate["fields_sha256"] in strings
            for decision in saved["decisions"]:
                assert decision["notes"] in text and decision["notes"] in strings
        for expected in (
            "unresolved",
            "Synthetic measurement: page 1, all gaps",
            "Kept for review",
        ):
            assert expected in text and expected in strings
        assert all(cell.data_type != "f" and cell.hyperlink is None for cell in cells)
        assert "must-never-export-source-json" not in text + strings
        assert str(case["source_path"]) not in text + strings
        assert len(pdf) < reports.MAX_OUTPUT_BYTES and len(xlsx) < reports.MAX_OUTPUT_BYTES
        db.commit()
    with case["factory"]() as db:
        actor = db.get(User, uid(100))
        assert reports.report_bytes(db, actor, case["draft_id"], report.id, "pdf") == pdf
        assert reports.report_bytes(db, actor, case["draft_id"], report.id, "xlsx") == xlsx
    book.close()
    (tmp_path / "candidate-review-report.pdf").write_bytes(pdf)
    (tmp_path / "candidate-review-report.xlsx").write_bytes(xlsx)


@pytest.mark.parametrize(
    "word,system", [(False, False), (False, True), (True, False), (True, True)]
)
def test_legacy_scope_and_singular_versions_keep_existing_output_layout(word, system):
    snapshot = report_snapshot(
        "scope-and-system" if system else "scope-only",
        word_scope() if word else _snapshot(imported=False)["scope"],
    )
    snapshot["render_version"] = (10 if system else 9) if word else (2 if system else 1)
    _seal(snapshot)
    reports.validate_report_snapshot(snapshot)
    assert "system_matches" not in snapshot
    assert outputs._collection_reviews(snapshot) == []
    book = load_workbook(io.BytesIO(outputs.render_scope_report_xlsx(snapshot)))
    assert "Selected reviews" not in book.sheetnames
    if system:
        assert [cell.value for cell in book["Review and coverage"][6]] == [
            "Part",
            "Field",
            "Saved value",
        ]
    text = _pdf_text(outputs.render_scope_report_pdf(snapshot))
    assert "Selected saved row reviews" not in text
    assert "review records are not quantities" not in text
    assert ("Saved technical review summary" in text) is system
    book.close()
