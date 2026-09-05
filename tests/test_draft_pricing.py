from __future__ import annotations

import copy
import hashlib
import io
import json
import zipfile

import pytest
import xlsxwriter
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pypdf import PdfReader
from sqlalchemy import func, select
from test_draft_estimates import add_payload
from test_draft_pdf_intake import pdf_setup as _pdf_setup
from test_draft_scope import sample_payload
from test_shared_file_containment import postgresql_session_factory as _postgres_fixture

from classifire.models import Estimate, LibraryRelease, StoredFile, User
from classifire.security import ROLE_PERMISSIONS
from classifire.services import draft_estimate_reports as reports
from classifire.services import draft_estimates as estimates
from classifire.services import draft_pricing_intake as pricing
from classifire.services import draft_scope as scope
from classifire.services.draft_pricing_contract import FIELDS, preview_rows
from classifire.services.draft_pricing_worker import process

pdf_setup = _pdf_setup
postgresql_session_factory = _postgres_fixture
MAPPING = {field: index for index, field in enumerate(FIELDS, 1)}


def workbook_bytes(*, rate=120.25, unit="each", currency="AUD", tax="excluded"):
    stream = io.BytesIO()
    book = xlsxwriter.Workbook(stream, {"in_memory": True, "strings_to_formulas": False})
    sheet = book.add_worksheet("Rates")
    sheet.write_row(0, 0, list(FIELDS))
    sheet.write_row(
        1,
        0,
        [
            "SYN-1",
            "Synthetic service work",
            unit,
            rate,
            currency,
            tax,
            "2026-09-06",
            40,
            60,
            "Service work",
            "Shared closure",
        ],
    )
    sheet.write_row(2, 0, ["UNKNOWN", "Awaiting price", "each", None, "AUD", "excluded"])
    sheet.write_row(3, 0, ["FORMULA", "Formula price", "each", None, "AUD", "excluded"])
    sheet.write_formula(3, 3, "=H2+I2", None, 100)
    book.close()
    return stream.getvalue()


def test_workbook_exact_cells_missing_and_formula_rates():
    content = workbook_bytes()
    document = json.loads(process(content))
    assert document["manifest"]["source_sha256"] == hashlib.sha256(content).hexdigest()
    rows = preview_rows(document, 1, 1, MAPPING)
    assert rows[0]["values"]["rate"] == "120.25"
    assert rows[0]["fields"]["rate"]["address"] == "D2"
    assert rows[0]["problems"] == []
    assert rows[1]["values"]["rate"] is None
    assert rows[1]["problems"]
    assert rows[2]["values"]["rate"] is None
    assert rows[2]["fields"]["rate"]["kind"] == "formula"
    assert rows[2]["fields"]["rate"]["value"] == "=H2+I2"
    assert rows[2]["problems"]
    with pytest.raises(ValueError, match="duplicate mapping"):
        preview_rows(document, 1, 1, {**MAPPING, "rate": 1})


@pytest.mark.parametrize(
    "changes,problem",
    [
        ({"rate": -1}, "rate_invalid"),
        ({"rate": "NaN"}, "rate_invalid"),
        ({"unit": "box"}, "unit_unsupported"),
        ({"currency": "USD"}, "currency_unsupported"),
        ({"tax": "included"}, "tax_basis_unsupported"),
    ],
)
def test_unsupported_prices_remain_unresolved(changes, problem):
    document = json.loads(process(workbook_bytes(**changes)))
    assert problem in preview_rows(document, 1, 1, MAPPING)[0]["problems"]


def test_workbook_selection_preserves_history_and_source_permissions(pdf_setup, monkeypatch):
    x = pdf_setup
    with x.factory() as db:
        actor = db.get(User, x.ids[0])
        scope.save_revision(db, actor, x.ids[2], 1, sample_payload())
        estimate = estimates.create_estimate(db, actor, x.ids[2], 2)
        envelope = estimates.add_line(db, actor, x.ids[2], estimate.id, 1, add_payload())
        line_id = envelope["lines"][0]["line_id"]
        original = estimates.revision_bytes(db, actor, x.ids[2], estimate.id, 2)
        source = pricing.intake().retain(
            db, actor, x.ids[2], "synthetic.xlsx", workbook_bytes(), settings=x.settings
        )
        db.commit()
        with pytest.raises(scope.DraftScopeError):
            pricing.preview(db, actor, x.ids[2], source.id, 1, 1, MAPPING, settings=x.settings)
        pricing.intake().scan_source(db, actor, x.ids[2], source.id, settings=x.settings)
        db.commit()
        _, rows = pricing.preview(
            db, actor, x.ids[2], source.id, 1, 1, MAPPING, settings=x.settings
        )

        def apply(row=rows[0], revision=2):
            return pricing.apply_rate(
                db,
                actor,
                x.ids[2],
                estimate.id,
                revision,
                line_id,
                source.id,
                1,
                1,
                MAPPING,
                row["row"],
                source.document_sha256,
                row["sha256"],
                "Service work only; shared closure excluded",
                settings=x.settings,
            )

        with pytest.raises(scope.DraftScopeError, match="PRICING_RATE_UNRESOLVED"):
            apply(rows[2])
        saved = apply()
        db.commit()
        assert saved["schema_version"] == "CLASSIFIRE-DRAFT-ESTIMATE-v2"
        assert saved["lines"][0]["unit_sell_rate"] == "120.25"
        assert saved["lines"][0]["original_rate"] == "1.005"
        assert saved["pricing_sources"][0]["row"]["fields"]["rate"]["address"] == "D2"
        assert estimates.revision_bytes(db, actor, x.ids[2], estimate.id, 2) == original
        with pytest.raises(scope.DraftScopeError):
            apply()
        assert db.scalar(select(func.count()).select_from(Estimate)) == 0
        assert db.scalar(select(func.count()).select_from(LibraryRelease)) == 0
        foreign = db.get(User, x.ids[1])
        with pytest.raises(scope.DraftScopeError) as denied:
            pricing.preview(db, foreign, x.ids[2], source.id, 1, 1, MAPPING, settings=x.settings)
        assert denied.value.status_code == 404
        tampered = copy.deepcopy(saved)
        tampered["pricing_sources"][0]["row"]["fields"]["rate"]["address"] = "D9"
        from classifire.services.draft_estimate_contract import validate_envelope

        with pytest.raises(ValueError):
            validate_envelope(tampered)
        report = reports.create_report(db, actor, x.ids[2], estimate.id, 3)
        pdf = reports.report_bytes(db, actor, x.ids[2], estimate.id, report.id, "pdf")
        xlsx = reports.report_bytes(db, actor, x.ids[2], estimate.id, report.id, "xlsx")
        assert "D2" in "".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf)).pages)
        book = load_workbook(io.BytesIO(xlsx), data_only=False)
        assert "Pricing sources" in book.sheetnames
        values = [str(cell.value) for row in book["Pricing sources"] for cell in row]
        assert "D2" in values and "120.25 (number)" in values
        assert not any(cell.data_type == "f" for row in book["Pricing sources"] for cell in row)
        book.close()
        reports.read_report(db, actor, x.ids[2], estimate.id, report.id)
        db.commit()
        assert (
            pricing.pricing_staleness(
                db, actor, x.ids[2], saved, storage_root=x.settings.storage_root
            )
            == []
        )
        stored = db.get(StoredFile, source.stored_file_id)
        stored.malware_scan_status = "malware_detected"
        db.commit()
        assert pricing.pricing_staleness(
            db, actor, x.ids[2], saved, storage_root=x.settings.storage_root
        ) == ["PRICING_SOURCE_UNAVAILABLE"]
        with pytest.raises(scope.DraftScopeError):
            apply(revision=3)
        with monkeypatch.context() as permissions:
            permissions.setitem(
                ROLE_PERMISSIONS, "estimator", ROLE_PERMISSIONS["estimator"] - {"library:read"}
            )
            for read in (
                lambda: estimates.read_estimate_revision(db, actor, x.ids[2], estimate.id),
                lambda: reports.report_bytes(db, actor, x.ids[2], estimate.id, report.id, "xlsx"),
            ):
                with pytest.raises(scope.DraftScopeError) as denied:
                    read()
                assert denied.value.status_code == 403
        estimate_id = estimate.id
    with x.factory() as db:
        actor = db.get(User, x.ids[0])
        assert estimates.read_estimate_revision(db, actor, x.ids[2], estimate_id) == saved


@pytest.mark.parametrize(
    "mode", ["hidden", "merged", "hyperlink", "oversized_grid", "duplicate", "macro", "external"]
)
def test_unsafe_or_unsupported_workbooks_fail_closed(mode):
    stream = io.BytesIO()
    book = xlsxwriter.Workbook(stream, {"in_memory": True})
    sheet = book.add_worksheet("Rates")
    sheet.write(0, 0, "Synthetic")
    if mode == "hidden":
        book.add_worksheet("Hidden").hide()
    elif mode == "merged":
        sheet.merge_range("A2:B2", "Merged")
    elif mode == "hyperlink":
        sheet.write_url("A2", "https://example.test/")
    elif mode == "oversized_grid":
        sheet.write(1000, 0, "Too far")
    book.close()
    content = stream.getvalue()
    if mode in ("duplicate", "macro", "external"):
        stream = io.BytesIO(content)
        with zipfile.ZipFile(stream, "a") as archive:
            name = {
                "duplicate": "XL/WORKBOOK.XML",
                "macro": "xl/vbaProject.bin",
                "external": "xl/externalLinks/externalLink1.xml",
            }[mode]
            archive.writestr(name, b"synthetic")
        content = stream.getvalue()
    with pytest.raises((ValueError, RuntimeError)):
        process(content)
