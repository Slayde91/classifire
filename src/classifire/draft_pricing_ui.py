"""Thin Draft pricing workbook preview and explicit rate-selection interface."""

from __future__ import annotations

from typing import Annotated, Any

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import HTMLResponse, RedirectResponse, Response
from openpyxl.utils.cell import get_column_letter  # type: ignore[import-untyped]

from .config import get_settings
from .draft_estimate_ui import Db, _actor, _revision
from .draft_scope_ui import _form_values
from .security import verify_csrf
from .services import draft_pricing_intake as pricing
from .services import malware_scan
from .services.draft_estimates import read_estimate_revision
from .services.draft_pricing_contract import (
    DATASET_KINDS,
    FIELDS,
    PRICE_MEANINGS,
    PROFILE_DECISIONS,
    ROW_EVIDENCE_STATES,
    ROW_ITEM_KINDS,
    SYSTEM_IDENTITY_EVIDENCE_FIELDS,
    SYSTEM_MAPPING_STATES,
    SYSTEM_MAPPING_UNRESOLVED_FIELDS,
)
from .services.draft_scope import DraftScopeError, get_draft
from .ui import _context, _require, templates
from .ui_uploads import single_file

router = APIRouter(include_in_schema=False)


async def _pricing_upload(request: Request, db: Db) -> dict[str, Any]:
    _require(request, db, "project:write")
    limit = min(get_settings().max_upload_bytes, malware_scan.MAX_SCAN_BYTES)
    return await single_file(request, limit, extra_fields=frozenset({"dataset_kind"}))


UploadData = Annotated[dict[str, Any], Depends(_pricing_upload)]


async def _pricing_form(request: Request) -> dict[str, str]:
    return await _form_values(request, 64 * 1024, max_fields=20)


FormData = Annotated[dict[str, str], Depends(_pricing_form)]


def _user(request: Request, db: Db, *, write: bool = False):
    user = _actor(request, db, write=write)
    _require(request, db, "library:read")
    return user


def _form_mapping(form: dict[str, str]) -> dict[str, int | None]:
    try:
        return {key: _revision(form[key]) if form[key] else None for key in FIELDS}
    except KeyError as exc:
        raise HTTPException(422, "Map the pricing fields explicitly") from exc


def _profile_revision(value: str) -> int:
    if not value.isascii() or not value.isdecimal() or not 1 <= len(value) <= 10:
        raise ValueError("profile revision")
    revision = int(value)
    if revision < 0:
        raise ValueError("profile revision")
    return revision


def _unresolved_fields(value: str) -> list[str]:
    fields = [] if not value.strip() else [item.strip() for item in value.split(",")]
    if fields != list(dict.fromkeys(fields)) or any(item not in FIELDS for item in fields):
        raise ValueError("unresolved fields")
    return fields


def _comma_values(
    value: str, *, allowed: tuple[str, ...] | None = None, maximum: int = 20
) -> list[str]:
    values = [] if not value.strip() else [item.strip() for item in value.split(",")]
    if (
        len(values) > maximum
        or values != list(dict.fromkeys(values))
        or any(not item or len(item) > 100 for item in values)
        or (allowed is not None and any(item not in allowed for item in values))
    ):
        raise ValueError("comma values")
    return values


def _page(
    request: Request,
    db: Db,
    user: Any,
    draft_id: str,
    estimate_id: str,
    source_id: str | None = None,
    profile_id: str | None = None,
    *,
    form: dict[str, str] | None = None,
    error: str | None = None,
    status_code: int = 200,
    row_observation_preview: dict[str, Any] | None = None,
    system_mapping_preview: dict[str, Any] | None = None,
) -> HTMLResponse:
    draft = get_draft(db, user, draft_id)
    envelope = read_estimate_revision(db, user, draft_id, estimate_id)
    source = None
    document = None
    rows: list[dict[str, Any]] = []
    profiles: list[dict[str, Any]] = []
    profile_decisions: list[dict[str, Any]] = []
    profile_preview = None
    selected_profile = None
    selected_profile_meta = None
    selected_decision = None
    row_observations: list[dict[str, Any]] = []
    system_mappings: list[dict[str, Any]] = []
    system_mapping_targets = None
    values = form or {}
    source_api = pricing.intake()
    if source_id is not None:
        source = pricing.source_info(db, user, draft_id, source_id)
        profiles = pricing.list_profiles(db, user, draft_id, source_id)
        profile_decisions = pricing.list_profile_decisions(db, user, draft_id, source_id)
        if profile_id is not None:
            selected_profile = pricing.read_profile(db, user, draft_id, source_id, profile_id)
            selected_profile_meta = next(item for item in profiles if item["id"] == profile_id)
            selected_decision = next(
                (item for item in profile_decisions if item["profile_id"] == profile_id), None
            )
            definition = selected_profile["definition"]
            if definition["dataset"]["kind"] == "general_pricelist":
                row_observations = pricing.list_row_observations(
                    db, user, draft_id, source_id, profile_id
                )
            else:
                try:
                    system_mappings = pricing.list_system_mappings(
                        db,
                        user,
                        draft_id,
                        source_id,
                        profile_id,
                        settings=get_settings(),
                    )
                    system_mapping_targets = pricing.list_system_mapping_targets(
                        db, user, draft_id
                    )
                except DraftScopeError as exc:
                    if exc.status_code != 403:
                        raise
            values = {
                "sheet_index": str(definition["selection"]["sheet_index"]),
                "header_row": str(definition["selection"]["header_row"]),
                "price_meaning": definition["commercial_basis"]["price_meaning"],
                **{
                    field: str(column) if column is not None else ""
                    for field, column in definition["selection"]["mapping"].items()
                },
            }
        if source["ready"]:
            try:
                _, document, _ = source_api._document(
                    db, user, draft_id, source_id, get_settings().storage_root
                )
                if form is not None and "price_meaning" in form:
                    profile_preview = pricing.preview_profile(
                        db,
                        user,
                        draft_id,
                        source_id,
                        _revision(form["sheet_index"]),
                        _revision(form["header_row"]),
                        _form_mapping(form),
                        form["price_meaning"],
                        settings=get_settings(),
                    )
                    rows = profile_preview["rows"]
                elif form is not None and "sheet_index" in form:
                    _, rows = pricing.preview(
                        db,
                        user,
                        draft_id,
                        source_id,
                        _revision(form["sheet_index"]),
                        _revision(form["header_row"]),
                        _form_mapping(form),
                        settings=get_settings(),
                    )
                elif selected_profile is not None:
                    definition = selected_profile["definition"]
                    _, rows = pricing.preview(
                        db,
                        user,
                        draft_id,
                        source_id,
                        definition["selection"]["sheet_index"],
                        definition["selection"]["header_row"],
                        definition["selection"]["mapping"],
                        settings=get_settings(),
                    )
            except DraftScopeError as exc:
                error = exc.code
                status_code = exc.status_code
    grids = []
    if document:
        for sheet in document["sheets"]:
            cells = {(cell["row"], cell["column"]): cell for cell in sheet["cells"]}
            grids.append(
                {
                    "name": sheet["name"],
                    "index": sheet["index"],
                    "rows": [
                        {
                            "number": number,
                            "cells": [
                                cells.get((number, col)) for col in range(1, sheet["columns"] + 1)
                            ],
                        }
                        for number in range(1, min(sheet["rows"], 12) + 1)
                    ],
                    "columns": [get_column_letter(col) for col in range(1, sheet["columns"] + 1)],
                }
            )
    base_url = f"/scopes/{draft_id}/estimates/{estimate_id}"
    return templates.TemplateResponse(
        request,
        "draft_pricing.html",
        _context(
            request,
            db,
            draft=draft,
            project=draft.project,
            envelope=envelope,
            estimate_url=base_url,
            pricing_url=base_url + "/pricing",
            source=source,
            sources=pricing.list_sources(db, user, draft_id),
            document=document,
            rows=rows,
            profiles=profiles,
            profile_decisions=profile_decisions,
            profile_preview=profile_preview,
            selected_profile=selected_profile,
            selected_profile_meta=selected_profile_meta,
            selected_decision=selected_decision,
            row_observations=row_observations,
            row_observation_preview=row_observation_preview,
            system_mappings=system_mappings,
            system_mapping_targets=system_mapping_targets,
            system_mapping_preview=system_mapping_preview,
            profile_decisions_allowed=PROFILE_DECISIONS,
            row_item_kinds=ROW_ITEM_KINDS,
            row_evidence_states=ROW_EVIDENCE_STATES,
            system_mapping_states=SYSTEM_MAPPING_STATES,
            system_identity_evidence_fields=SYSTEM_IDENTITY_EVIDENCE_FIELDS,
            system_mapping_unresolved_fields=SYSTEM_MAPPING_UNRESOLVED_FIELDS,
            fields=FIELDS,
            dataset_kinds=DATASET_KINDS,
            price_meanings=PRICE_MEANINGS,
            grids=grids,
            columns=[(col, get_column_letter(col)) for col in range(1, 51)],
            form_values=values,
            error=error,
            postgres=db.get_bind().dialect.name == "postgresql",
        ),
        status_code=status_code,
        headers={"Cache-Control": "no-store"},
    )


@router.get("/scopes/{draft_id}/estimates/{estimate_id}/pricing", response_class=HTMLResponse)
def pricing_page(request: Request, db: Db, draft_id: str, estimate_id: str) -> HTMLResponse:
    user = _user(request, db)
    try:
        return _page(request, db, user, draft_id, estimate_id)
    except DraftScopeError as exc:
        raise HTTPException(exc.status_code, exc.code) from exc


@router.post("/scopes/{draft_id}/estimates/{estimate_id}/pricing/upload")
def upload_pricing(
    request: Request, db: Db, draft_id: str, estimate_id: str, data: UploadData
) -> RedirectResponse:
    user = _user(request, db, write=True)
    try:
        read_estimate_revision(db, user, draft_id, estimate_id)
        source = pricing.retain_source(
            db,
            user,
            draft_id,
            data["filename"],
            data["content"],
            data["dataset_kind"],
            settings=get_settings(),
        )
        db.commit()
    except DraftScopeError as exc:
        db.rollback()
        raise HTTPException(exc.status_code, exc.code) from exc
    return RedirectResponse(f"/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source.id}", 303)


@router.get(
    "/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source_id}", response_class=HTMLResponse
)
def pricing_source(
    request: Request, db: Db, draft_id: str, estimate_id: str, source_id: str
) -> HTMLResponse:
    user = _user(request, db)
    try:
        return _page(request, db, user, draft_id, estimate_id, source_id)
    except DraftScopeError as exc:
        raise HTTPException(exc.status_code, exc.code) from exc


@router.post("/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source_id}/scan")
def scan_pricing(
    request: Request, db: Db, draft_id: str, estimate_id: str, source_id: str, form: FormData
) -> RedirectResponse:
    verify_csrf(request, form.get("csrf_token"))
    user = _user(request, db, write=True)
    if set(form) != {"csrf_token"}:
        raise HTTPException(422, "Submit one explicit scan request")
    try:
        read_estimate_revision(db, user, draft_id, estimate_id)
        pricing.intake().scan_source(db, user, draft_id, source_id, settings=get_settings())
        db.commit()
    except DraftScopeError as exc:
        db.rollback()
        raise HTTPException(exc.status_code, exc.code) from exc
    return RedirectResponse(f"/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source_id}", 303)


@router.post(
    "/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source_id}/preview",
    response_class=HTMLResponse,
)
def preview_pricing(
    request: Request, db: Db, draft_id: str, estimate_id: str, source_id: str, form: FormData
) -> HTMLResponse:
    verify_csrf(request, form.get("csrf_token"))
    user = _user(request, db)
    if set(form) != {"csrf_token", "sheet_index", "header_row", "price_meaning", *FIELDS}:
        raise HTTPException(422, "Map only the supported pricing fields")
    try:
        return _page(request, db, user, draft_id, estimate_id, source_id, form=form)
    except DraftScopeError as exc:
        db.rollback()
        raise HTTPException(exc.status_code, exc.code) from exc


@router.post("/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source_id}/profiles")
def save_pricing_profile(
    request: Request, db: Db, draft_id: str, estimate_id: str, source_id: str, form: FormData
) -> RedirectResponse:
    verify_csrf(request, form.get("csrf_token"))
    user = _user(request, db, write=True)
    if set(form) != {
        "csrf_token",
        "sheet_index",
        "header_row",
        "price_meaning",
        *FIELDS,
        "expected_profile_revision",
        "document_sha256",
        "preview_hash",
    }:
        raise HTTPException(422, "Confirm only the previewed pricing source profile")
    try:
        read_estimate_revision(db, user, draft_id, estimate_id)
        saved = pricing.save_profile(
            db,
            user,
            draft_id,
            source_id,
            _revision(form["sheet_index"]),
            _revision(form["header_row"]),
            _form_mapping(form),
            form["price_meaning"],
            _profile_revision(form["expected_profile_revision"]),
            form["document_sha256"],
            form["preview_hash"],
            settings=get_settings(),
        )
        db.commit()
    except (DraftScopeError, ValueError) as exc:
        db.rollback()
        if isinstance(exc, DraftScopeError):
            raise HTTPException(exc.status_code, exc.code) from exc
        raise HTTPException(422, "Choose a valid profile revision") from exc
    return RedirectResponse(
        f"/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source_id}/profiles/"
        + saved["profile_id"],
        303,
    )


@router.get(
    "/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source_id}/profiles/{profile_id}",
    response_class=HTMLResponse,
)
def pricing_profile(
    request: Request,
    db: Db,
    draft_id: str,
    estimate_id: str,
    source_id: str,
    profile_id: str,
) -> HTMLResponse:
    user = _user(request, db)
    try:
        return _page(
            request,
            db,
            user,
            draft_id,
            estimate_id,
            source_id,
            profile_id=profile_id,
        )
    except DraftScopeError as exc:
        raise HTTPException(exc.status_code, exc.code) from exc


@router.get(
    "/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source_id}/profiles/"
    "{profile_id}/download"
)
def download_pricing_profile(
    request: Request,
    db: Db,
    draft_id: str,
    estimate_id: str,
    source_id: str,
    profile_id: str,
) -> Response:
    user = _user(request, db)
    try:
        read_estimate_revision(db, user, draft_id, estimate_id)
        content = pricing.profile_bytes(db, user, draft_id, source_id, profile_id)
    except DraftScopeError as exc:
        raise HTTPException(exc.status_code, exc.code) from exc
    return Response(
        content,
        media_type="application/json",
        headers={
            "Content-Disposition": f'attachment; filename="pricing-profile-{profile_id}.json"',
            "Cache-Control": "no-store",
            "X-Content-Type-Options": "nosniff",
        },
    )


@router.post(
    "/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source_id}/profiles/"
    "{profile_id}/decisions"
)
def review_pricing_profile(
    request: Request,
    db: Db,
    draft_id: str,
    estimate_id: str,
    source_id: str,
    profile_id: str,
    form: FormData,
) -> RedirectResponse:
    verify_csrf(request, form.get("csrf_token"))
    user = _user(request, db)
    _require(request, db, "pricing:approve")
    if set(form) != {"csrf_token", "profile_sha256", "decision", "reason"}:
        raise HTTPException(422, "Review only the selected pricing source profile")
    try:
        read_estimate_revision(db, user, draft_id, estimate_id)
        saved = pricing.save_profile_decision(
            db,
            user,
            draft_id,
            source_id,
            profile_id,
            form["profile_sha256"],
            form["decision"],
            form["reason"],
        )
        db.commit()
    except DraftScopeError as exc:
        db.rollback()
        raise HTTPException(exc.status_code, exc.code) from exc
    return RedirectResponse(
        f"/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source_id}/profiles/"
        + saved["profile_id"],
        303,
    )


@router.get(
    "/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source_id}/profiles/"
    "{profile_id}/decisions/{decision_id}/download"
)
def download_pricing_profile_decision(
    request: Request,
    db: Db,
    draft_id: str,
    estimate_id: str,
    source_id: str,
    profile_id: str,
    decision_id: str,
) -> Response:
    user = _user(request, db)
    try:
        read_estimate_revision(db, user, draft_id, estimate_id)
        content = pricing.profile_decision_bytes(
            db, user, draft_id, source_id, profile_id, decision_id
        )
    except DraftScopeError as exc:
        raise HTTPException(exc.status_code, exc.code) from exc
    return Response(
        content,
        media_type="application/json",
        headers={
            "Content-Disposition": (
                f'attachment; filename="pricing-profile-decision-{decision_id}.json"'
            ),
            "Cache-Control": "no-store",
            "X-Content-Type-Options": "nosniff",
        },
    )


@router.post(
    "/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source_id}/profiles/"
    "{profile_id}/rows/{row_number}/preview",
    response_class=HTMLResponse,
)
def preview_pricing_row_observation(
    request: Request,
    db: Db,
    draft_id: str,
    estimate_id: str,
    source_id: str,
    profile_id: str,
    row_number: int,
    form: FormData,
) -> HTMLResponse:
    verify_csrf(request, form.get("csrf_token"))
    user = _user(request, db)
    _require(request, db, "pricing:approve")
    if set(form) != {
        "csrf_token",
        "item_kind",
        "normalized_reference",
        "evidence_state",
        "review_reason",
        "unresolved_fields",
    }:
        raise HTTPException(422, "Review only the selected exact pricing row")
    try:
        read_estimate_revision(db, user, draft_id, estimate_id)
        preview = pricing.preview_row_observation(
            db,
            user,
            draft_id,
            source_id,
            profile_id,
            row_number,
            form["item_kind"],
            form["normalized_reference"],
            form["evidence_state"],
            form["review_reason"],
            _unresolved_fields(form["unresolved_fields"]),
            settings=get_settings(),
        )
        return _page(
            request,
            db,
            user,
            draft_id,
            estimate_id,
            source_id,
            profile_id,
            form=form,
            row_observation_preview=preview,
        )
    except DraftScopeError as exc:
        raise HTTPException(exc.status_code, exc.code) from exc
    except ValueError as exc:
        raise HTTPException(422, "Choose a valid row interpretation") from exc


@router.post(
    "/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source_id}/profiles/"
    "{profile_id}/rows/{row_number}/observations"
)
def save_pricing_row_observation(
    request: Request,
    db: Db,
    draft_id: str,
    estimate_id: str,
    source_id: str,
    profile_id: str,
    row_number: int,
    form: FormData,
) -> RedirectResponse:
    verify_csrf(request, form.get("csrf_token"))
    user = _user(request, db)
    _require(request, db, "pricing:approve")
    if set(form) != {
        "csrf_token",
        "item_kind",
        "normalized_reference",
        "evidence_state",
        "review_reason",
        "unresolved_fields",
        "profile_sha256",
        "decision_sha256",
        "row_sha256",
        "preview_hash",
    }:
        raise HTTPException(422, "Save only the previewed exact pricing row")
    try:
        read_estimate_revision(db, user, draft_id, estimate_id)
        pricing.save_row_observation(
            db,
            user,
            draft_id,
            source_id,
            profile_id,
            row_number,
            form["item_kind"],
            form["normalized_reference"],
            form["evidence_state"],
            form["review_reason"],
            _unresolved_fields(form["unresolved_fields"]),
            form["profile_sha256"],
            form["decision_sha256"],
            form["row_sha256"],
            form["preview_hash"],
            settings=get_settings(),
        )
        db.commit()
    except DraftScopeError as exc:
        db.rollback()
        raise HTTPException(exc.status_code, exc.code) from exc
    except ValueError as exc:
        db.rollback()
        raise HTTPException(422, "Choose a valid row interpretation") from exc
    return RedirectResponse(
        f"/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source_id}/profiles/{profile_id}",
        303,
    )


@router.get(
    "/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source_id}/profiles/"
    "{profile_id}/observations/{observation_id}/download"
)
def download_pricing_row_observation(
    request: Request,
    db: Db,
    draft_id: str,
    estimate_id: str,
    source_id: str,
    profile_id: str,
    observation_id: str,
) -> Response:
    user = _user(request, db)
    try:
        read_estimate_revision(db, user, draft_id, estimate_id)
        content = pricing.row_observation_bytes(
            db, user, draft_id, source_id, profile_id, observation_id
        )
    except DraftScopeError as exc:
        raise HTTPException(exc.status_code, exc.code) from exc
    return Response(
        content,
        media_type="application/json",
        headers={
            "Content-Disposition": (
                f'attachment; filename="pricing-row-observation-{observation_id}.json"'
            ),
            "Cache-Control": "no-store",
            "X-Content-Type-Options": "nosniff",
        },
    )


@router.post(
    "/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source_id}/profiles/"
    "{profile_id}/rows/{row_number}/system-mapping/preview",
    response_class=HTMLResponse,
)
def preview_pricing_system_mapping(
    request: Request,
    db: Db,
    draft_id: str,
    estimate_id: str,
    source_id: str,
    profile_id: str,
    row_number: int,
    form: FormData,
) -> HTMLResponse:
    verify_csrf(request, form.get("csrf_token"))
    user = _user(request, db)
    _require(request, db, "pricing:approve")
    _require(request, db, "technical:read")
    if set(form) != {
        "csrf_token",
        "technical_release_id",
        "mapping_status",
        "normalized_reference",
        "selected_variant_id",
        "candidate_variant_ids",
        "identity_evidence_fields",
        "review_reason",
        "unresolved_fields",
    }:
        raise HTTPException(422, "Map only the selected exact Firefly system-price row")
    try:
        read_estimate_revision(db, user, draft_id, estimate_id)
        candidates = sorted(_comma_values(form["candidate_variant_ids"]))
        preview = pricing.preview_system_mapping(
            db,
            user,
            draft_id,
            source_id,
            profile_id,
            row_number,
            form["technical_release_id"],
            form["mapping_status"],
            form["normalized_reference"],
            form["selected_variant_id"] or None,
            candidates,
            _comma_values(
                form["identity_evidence_fields"],
                allowed=SYSTEM_IDENTITY_EVIDENCE_FIELDS,
            ),
            form["review_reason"],
            _comma_values(
                form["unresolved_fields"],
                allowed=SYSTEM_MAPPING_UNRESOLVED_FIELDS,
            ),
            settings=get_settings(),
        )
        return _page(
            request,
            db,
            user,
            draft_id,
            estimate_id,
            source_id,
            profile_id,
            form=form,
            system_mapping_preview=preview,
        )
    except DraftScopeError as exc:
        raise HTTPException(exc.status_code, exc.code) from exc
    except ValueError as exc:
        raise HTTPException(422, "Choose a valid system mapping") from exc


@router.post(
    "/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source_id}/profiles/"
    "{profile_id}/rows/{row_number}/system-mappings"
)
def save_pricing_system_mapping(
    request: Request,
    db: Db,
    draft_id: str,
    estimate_id: str,
    source_id: str,
    profile_id: str,
    row_number: int,
    form: FormData,
) -> RedirectResponse:
    verify_csrf(request, form.get("csrf_token"))
    user = _user(request, db)
    _require(request, db, "pricing:approve")
    _require(request, db, "technical:read")
    if set(form) != {
        "csrf_token",
        "technical_release_id",
        "mapping_status",
        "normalized_reference",
        "selected_variant_id",
        "candidate_variant_ids",
        "identity_evidence_fields",
        "review_reason",
        "unresolved_fields",
        "profile_sha256",
        "decision_sha256",
        "row_sha256",
        "technical_release_sha256",
        "preview_hash",
    }:
        raise HTTPException(422, "Save only the previewed Firefly system-price mapping")
    try:
        read_estimate_revision(db, user, draft_id, estimate_id)
        pricing.save_system_mapping(
            db,
            user,
            draft_id,
            source_id,
            profile_id,
            row_number,
            form["technical_release_id"],
            form["mapping_status"],
            form["normalized_reference"],
            form["selected_variant_id"] or None,
            sorted(_comma_values(form["candidate_variant_ids"])),
            _comma_values(
                form["identity_evidence_fields"],
                allowed=SYSTEM_IDENTITY_EVIDENCE_FIELDS,
            ),
            form["review_reason"],
            _comma_values(
                form["unresolved_fields"],
                allowed=SYSTEM_MAPPING_UNRESOLVED_FIELDS,
            ),
            form["profile_sha256"],
            form["decision_sha256"],
            form["row_sha256"],
            form["technical_release_sha256"],
            form["preview_hash"],
            settings=get_settings(),
        )
        db.commit()
    except DraftScopeError as exc:
        db.rollback()
        raise HTTPException(exc.status_code, exc.code) from exc
    except ValueError as exc:
        db.rollback()
        raise HTTPException(422, "Choose a valid system mapping") from exc
    return RedirectResponse(
        f"/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source_id}/profiles/{profile_id}",
        303,
    )


@router.get(
    "/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source_id}/profiles/"
    "{profile_id}/system-mappings/{mapping_id}/download"
)
def download_pricing_system_mapping(
    request: Request,
    db: Db,
    draft_id: str,
    estimate_id: str,
    source_id: str,
    profile_id: str,
    mapping_id: str,
) -> Response:
    user = _user(request, db)
    _require(request, db, "pricing:approve")
    _require(request, db, "technical:read")
    try:
        read_estimate_revision(db, user, draft_id, estimate_id)
        content = pricing.system_mapping_bytes(
            db, user, draft_id, source_id, profile_id, mapping_id
        )
    except DraftScopeError as exc:
        raise HTTPException(exc.status_code, exc.code) from exc
    return Response(
        content,
        media_type="application/json",
        headers={
            "Content-Disposition": (
                f'attachment; filename="pricing-system-mapping-{mapping_id}.json"'
            ),
            "Cache-Control": "no-store",
            "X-Content-Type-Options": "nosniff",
        },
    )


@router.post(
    "/scopes/{draft_id}/estimates/{estimate_id}/pricing/{source_id}/apply", response_model=None
)
def apply_pricing(
    request: Request, db: Db, draft_id: str, estimate_id: str, source_id: str, form: FormData
) -> HTMLResponse | RedirectResponse:
    verify_csrf(request, form.get("csrf_token"))
    user = _user(request, db, write=True)
    if set(form) != {
        "csrf_token",
        "sheet_index",
        "header_row",
        *FIELDS,
        "line_id",
        "expected_revision",
        "row_number",
        "document_sha256",
        "row_sha256",
        "recovery_note",
    }:
        raise HTTPException(422, "Select one previewed row and explain its recovery scope")
    try:
        pricing.apply_rate(
            db,
            user,
            draft_id,
            estimate_id,
            _revision(form["expected_revision"]),
            form["line_id"],
            source_id,
            _revision(form["sheet_index"]),
            _revision(form["header_row"]),
            _form_mapping(form),
            _revision(form["row_number"]),
            form["document_sha256"],
            form["row_sha256"],
            form["recovery_note"],
            settings=get_settings(),
        )
        db.commit()
    except DraftScopeError as exc:
        db.rollback()
        if exc.status_code not in {409, 422}:
            raise HTTPException(exc.status_code, exc.code) from exc
        return _page(
            request,
            db,
            user,
            draft_id,
            estimate_id,
            source_id,
            form=form,
            error="No rate was applied. " + exc.code,
            status_code=exc.status_code,
        )
    return RedirectResponse(f"/scopes/{draft_id}/estimates/{estimate_id}", 303)
