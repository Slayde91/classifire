from __future__ import annotations

import copy
import hashlib
import io
from decimal import Decimal
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pypdf import PdfReader
from sqlalchemy import func, select, update
from test_draft_estimates import add_payload, prepare
from test_draft_scope import sample_payload, setup, uid  # noqa: F401
from test_draft_system_matches import case as match_case  # noqa: F401
from test_draft_system_matches import create as create_match

from classifire.models import AuditEvent, DraftEstimateReport, Estimate, EstimateLine, User
from classifire.outputs import draft_estimate as outputs
from classifire.services import draft_estimate_reports as reports
from classifire.services import draft_estimates as estimates
from classifire.services.draft_scope import DraftScopeError, save_revision


def priced(db, **changes):
    actor, draft, estimate = prepare(db)
    envelope = estimates.add_line(db, actor, draft.id, estimate.id, 1, add_payload(**changes))
    db.commit()
    return actor, draft, estimate, envelope


def test_real_report_pair_preserves_snapshot_and_survives_fresh_session(setup, tmp_path):  # noqa: F811
    with setup() as db:
        actor, draft, estimate, envelope = priced(db)
        row = reports.create_report(db, actor, draft.id, estimate.id, 2)
        snapshot = reports.read_report(db, actor, draft.id, estimate.id, row.id)
        pdf = reports.report_bytes(db, actor, draft.id, estimate.id, row.id, "pdf")
        xlsx = reports.report_bytes(db, actor, draft.id, estimate.id, row.id, "xlsx")
        assert snapshot["estimate"] == envelope
        assert snapshot["profile"] == "estimate-only"
        text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf)).pages)
        assert "2.01" in text and "1.005" in text and "Original quantity: 2" in text
        assert "Tax: Not calculated" in text and "Site verification required" in text
        book = load_workbook(io.BytesIO(xlsx), data_only=False)
        assert book.sheetnames == [
            "Summary",
            "Lines",
            "Line details",
            "History",
            "Coverage",
            "Scope context",
        ]
        cells = list(book["Lines"].iter_rows(min_row=7, max_row=7))[0]
        assert [cell.value for cell in cells[3:9]] == ["each", "2", "1.005", "2.01", 2.01, "priced"]
        assert cells[6].data_type == "s" and cells[7].data_type == "n"
        for sheet in book:
            assert sheet.freeze_panes == "C7"
            assert sheet.auto_filter.ref
            assert not any(cell.data_type == "f" or cell.hyperlink for row in sheet for cell in row)
        with ZipFile(io.BytesIO(xlsx)) as archive:
            assert archive.read("xl/media/image1.png") == outputs._logo_path().read_bytes()
        old_hashes = (row.pdf_sha256, row.xlsx_sha256)
        ids = (draft.id, estimate.id, row.id)
        db.commit()
        newer = estimates.override_line(
            db,
            actor,
            draft.id,
            estimate.id,
            2,
            envelope["lines"][0]["line_id"],
            {
                "quantity": "3",
                "unit_sell_rate": "2",
                "reason": "Synthetic later input",
            },
        )
        assert newer["summary"]["priced_subtotal_ex_tax"] == "6.00"
        save_revision(db, actor, draft.id, 2, sample_payload())
        draft.project.name = "Changed project title"
        db.commit()
        stale = reports.report_staleness(db, actor, *ids, storage_root=tmp_path)
        assert {
            "REPORT_ESTIMATE_CHANGED",
            "ESTIMATE_SCOPE_CHANGED",
            "REPORT_PROJECT_CHANGED",
        } <= set(stale)
    with setup() as db:
        actor = db.get(User, uid(100))
        assert reports.report_bytes(db, actor, *ids, "pdf") == pdf
        assert reports.report_bytes(db, actor, *ids, "xlsx") == xlsx
        assert [item.id for item in reports.list_reports(db, actor, ids[0], ids[1])] == [ids[2]]
        assert old_hashes == (hashlib.sha256(pdf).hexdigest(), hashlib.sha256(xlsx).hexdigest())
        assert db.scalar(select(func.count()).select_from(Estimate)) == 0
        assert db.scalar(select(func.count()).select_from(EstimateLine)) == 0


@pytest.mark.parametrize("value", ["0.00", "1.23", "9999999999999.99", "999999998000000001.00"])
def test_numeric_convenience_never_changes_authoritative_decimal(value):
    numeric = outputs._safe_numeric(value)
    if numeric is not None:
        assert Decimal(str(numeric)) == Decimal(value)
        assert len(Decimal(value).normalize().as_tuple().digits) <= 15
    if value == "999999998000000001.00":
        assert numeric is None


def test_large_amount_and_hostile_long_text_are_retained_exactly(setup):  # noqa: F811
    hostile = '=HYPERLINK("https://example.invalid", "run") <b>Literal</b> '
    note = hostile + "Long synthetic source note. " * 120 + "END-NOTE"
    with setup() as db:
        actor, draft, estimate, _ = priced(
            db,
            quantity="999999999",
            unit_sell_rate="999999999",
            reason="Explicit synthetic count",
            source_note=note,
        )
        row = reports.create_report(db, actor, draft.id, estimate.id, 2)
        data = reports.report_bytes(db, actor, draft.id, estimate.id, row.id, "xlsx")
        book = load_workbook(io.BytesIO(data), data_only=False)
        assert book["Lines"]["G7"].value == "999999998000000001.00"
        assert book["Lines"]["H7"].value is None
        restored = "".join(
            row[3].value or ""
            for row in book["Line details"].iter_rows(min_row=7)
            if row[2].value == "source note" or (row[2].value is None and row[3].value)
        )
        assert note in restored
        assert not any(
            cell.data_type == "f" or cell.hyperlink
            for sheet in book
            for row in sheet
            for cell in row
        )
        pdf = reports.report_bytes(db, actor, draft.id, estimate.id, row.id, "pdf")
        text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf)).pages)
        assert "999999998000000001.00" in text and "END-NOTE" in text and "<b>Literal</b>" in text


@pytest.mark.parametrize("field", ["pdf_bytes", "xlsx_bytes", "snapshot_json", "estimate_hash"])
def test_corruption_refuses_both_downloads_without_audit(setup, field):  # noqa: F811
    with setup() as db:
        actor, draft, estimate, _ = priced(db)
        row = reports.create_report(db, actor, draft.id, estimate.id, 2)
        db.commit()
        value = b"corrupt" if field.endswith("bytes") else "corrupt"
        db.execute(
            update(DraftEstimateReport)
            .where(DraftEstimateReport.id == row.id)
            .values(**{field: value})
        )
        db.commit()
        count = db.scalar(select(func.count()).select_from(AuditEvent))
        for fmt in ("pdf", "xlsx"):
            with pytest.raises(DraftScopeError) as error:
                reports.report_bytes(db, actor, draft.id, estimate.id, row.id, fmt)
            assert error.value.status_code == 409
        assert db.scalar(select(func.count()).select_from(AuditEvent)) == count


@pytest.mark.parametrize("failure", ["second_output", "mutation", "revocation"])
def test_generation_is_atomic_and_rechecks_authority(setup, monkeypatch, failure):  # noqa: F811
    with setup() as db:
        actor, draft, estimate, _ = priced(db)
        before = db.scalar(select(func.count()).select_from(AuditEvent))

        def render(snapshot):
            if failure == "mutation":
                snapshot["project"]["name"] = "Mutated"
            if failure == "revocation":
                db.execute(update(User).where(User.id == actor.id).values(is_active=False))
                db.flush()
            return b"%PDF-synthetic"

        def xlsx(snapshot):
            if failure == "second_output":
                raise ValueError("private renderer detail must not enter safe failure")
            return b"PK\x03\x04synthetic"

        monkeypatch.setattr(outputs, "render_estimate_report_pdf", render)
        monkeypatch.setattr(outputs, "render_estimate_report_xlsx", xlsx)
        with pytest.raises(DraftScopeError) as error:
            reports.create_report(db, actor, draft.id, estimate.id, 2)
        assert "private" not in str(error.value)
        assert db.scalar(select(func.count()).select_from(DraftEstimateReport)) == 0
        assert db.scalar(select(func.count()).select_from(AuditEvent)) == before


def test_foreign_owner_and_revoked_export_are_denied(setup, monkeypatch):  # noqa: F811
    with setup() as db:
        actor, draft, estimate, _ = priced(db)
        row = reports.create_report(db, actor, draft.id, estimate.id, 2)
        db.commit()
        foreign = db.get(User, uid(101))
        with pytest.raises(DraftScopeError):
            reports.read_report(db, foreign, draft.id, estimate.id, row.id)
        from classifire.services import draft_scope

        original = draft_scope._actor

        def restricted(db, user, permission):
            if permission == "estimate:export":
                raise DraftScopeError("PERMISSION_DENIED", 403)
            return original(db, user, permission)

        monkeypatch.setattr("classifire.services.draft_estimates._actor", restricted)
        assert (
            reports.read_report(db, actor, draft.id, estimate.id, row.id)["estimate"]["revision"]
            == 2
        )
        with pytest.raises(DraftScopeError) as error:
            reports.report_bytes(db, actor, draft.id, estimate.id, row.id, "pdf")
        assert error.value.status_code == 403


def test_pure_snapshot_refuses_resealed_false_amount(setup):  # noqa: F811
    with setup() as db:
        actor, draft, estimate, _ = priced(db)
        row = reports.create_report(db, actor, draft.id, estimate.id, 2)
        snapshot = copy.deepcopy(reports.read_report(db, actor, draft.id, estimate.id, row.id))
        snapshot["estimate"]["summary"]["priced_subtotal_ex_tax"] = "999.99"
        snapshot["sha256"] = reports._checksum(snapshot)
        with pytest.raises(DraftScopeError):
            outputs.render_estimate_report_pdf(snapshot)


@pytest.mark.parametrize(
    ("quantity", "rate", "status", "amount"),
    [
        (None, "0", "active", None),
        ("0", None, "active", None),
        ("0", "1", "active", "0.00"),
        ("2", "1", "omitted", None),
    ],
)
def test_unknown_zero_and_omitted_values_remain_distinct(setup, quantity, rate, status, amount):  # noqa: F811
    with setup() as db:
        actor, draft, estimate, envelope = priced(
            db, quantity=quantity, unit_sell_rate=rate, reason="Explicit synthetic quantity basis"
        )
        revision = 2
        if status == "omitted":
            estimates.set_line_status(
                db,
                actor,
                draft.id,
                estimate.id,
                2,
                envelope["lines"][0]["line_id"],
                "omitted",
                "Separate work",
            )
            revision = 3
        row = reports.create_report(db, actor, draft.id, estimate.id, revision)
        book = load_workbook(
            io.BytesIO(reports.report_bytes(db, actor, draft.id, estimate.id, row.id, "xlsx"))
        )
        assert book["Lines"]["E7"].value == (
            "Unknown / unavailable" if quantity is None else quantity
        )
        assert book["Lines"]["G7"].value == ("Unknown / unavailable" if amount is None else amount)
        assert book["Lines"]["I7"].value == (
            "omitted" if status == "omitted" else "unpriced" if amount is None else "priced"
        )


def test_attached_review_source_staleness_preserves_reports_and_requires_access(match_case):  # noqa: F811
    case = match_case
    with case["factory"]() as db:
        actor = db.get(User, uid(100))
        match = create_match(db, case)
        estimate = estimates.create_estimate(
            db, actor, case["draft_id"], 2, match_id=match.id, match_revision=1
        )
        row = reports.create_report(db, actor, case["draft_id"], estimate.id, 1)
        db.commit()
        ids = (case["draft_id"], estimate.id, row.id)
        original = reports.report_bytes(db, actor, *ids, "pdf")
        assert reports.report_staleness(db, actor, *ids, storage_root=case["storage_root"]) == []
        case["source_path"].write_bytes(b"Synthetic changed source bytes")
        assert reports.report_staleness(db, actor, *ids, storage_root=case["storage_root"])
        assert reports.report_bytes(db, actor, *ids, "pdf") == original
        actor.role = "project_manager"
        db.commit()
        with pytest.raises(DraftScopeError) as error:
            reports.read_report(db, actor, *ids)
        assert error.value.status_code == 403


def test_export_permission_rechecked_after_output_integrity_work(setup, monkeypatch):  # noqa: F811
    with setup() as db:
        actor, draft, estimate, _ = priced(db)
        row = reports.create_report(db, actor, draft.id, estimate.id, 2)
        db.commit()
        original = reports._output

        def output(value, fmt):
            content = original(value, fmt)
            db.execute(update(User).where(User.id == actor.id).values(role="read_only"))
            db.flush()
            return content

        monkeypatch.setattr(reports, "_output", output)
        before = db.scalar(select(func.count()).select_from(AuditEvent))
        with pytest.raises(DraftScopeError) as error:
            reports.report_bytes(db, actor, draft.id, estimate.id, row.id, "pdf")
        assert error.value.status_code == 403
        assert db.scalar(select(func.count()).select_from(AuditEvent)) == before
