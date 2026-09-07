"""Explicit workbook mapping and rate provenance, with no inferred prices or authority."""

from __future__ import annotations

import re
from datetime import UTC, date, datetime
from typing import Any

from .draft_system_match_contract import canonical, digest, validate_binding
from .technical_field_snapshot import FIELD_NAMES as TECHNICAL_FIELD_NAMES

FIELDS = (
    "reference",
    "description",
    "unit",
    "rate",
    "currency",
    "tax_basis",
    "rate_date",
    "labour",
    "materials",
    "inclusions",
    "exclusions",
)
REQUIRED_MAPPING = ("reference", "description", "unit", "rate", "currency", "tax_basis")
PROFILE_SCHEMA = "CLASSIFIRE-DRAFT-PRICING-SOURCE-PROFILE-v1"
PROFILE_DECISION_SCHEMA = "CLASSIFIRE-DRAFT-PRICING-SOURCE-PROFILE-DECISION-v1"
ROW_OBSERVATION_SCHEMA = "CLASSIFIRE-DRAFT-PRICING-ROW-OBSERVATION-v1"
SYSTEM_MAPPING_SCHEMA = "CLASSIFIRE-DRAFT-PRICING-SYSTEM-MAPPING-v1"
PROFILE_DECISIONS = ("approve", "reject", "request_revision")
ROW_ITEM_KINDS = ("product", "material", "labour", "service")
ROW_EVIDENCE_STATES = ("confirmed", "provisional")
SYSTEM_MAPPING_STATES = ("mapped", "unmatched", "ambiguous")
SYSTEM_IDENTITY_EVIDENCE_FIELDS = ("reference", "description", "inclusions", "exclusions")
SYSTEM_MAPPING_UNRESOLVED_FIELDS = (
    *FIELDS,
    "technical_variant",
    "configuration_identity",
    "source_alias",
    "commercial_scope",
)
DATASET_KINDS = ("general_pricelist", "firefly_system_prices")
PRICE_MEANINGS = (
    "unknown",
    "buy_cost",
    "list_price",
    "sell_price",
    "quoted_price",
    "actual_price",
)


def validate_mapping(mapping: Any, columns: int) -> None:
    if type(mapping) is not dict or set(mapping) != set(FIELDS):
        raise ValueError("mapping")
    used = []
    for key, value in mapping.items():
        if value is None and key not in REQUIRED_MAPPING:
            continue
        if type(value) is not int or not 1 <= value <= columns:
            raise ValueError("mapping column")
        used.append(value)
    if len(set(used)) != len(used):
        raise ValueError("duplicate mapping")


def row_values(fields: dict[str, Any]) -> tuple[dict[str, str | None], list[str]]:
    from .draft_estimate_contract import decimal_string

    values = {key: cell["value"] if cell is not None else None for key, cell in fields.items()}
    problems = []
    for key in REQUIRED_MAPPING:
        cell = fields[key]
        if cell is None or cell["kind"] not in ("text", "number") or not values[key]:
            problems.append(key + "_missing_or_unsupported")
    rate = fields["rate"]
    if rate is not None and rate["kind"] in ("text", "number"):
        try:
            values["rate"] = decimal_string(rate["value"])
        except (ValueError, ArithmeticError):
            problems.append("rate_invalid")
    else:
        values["rate"] = None
    if values["unit"] not in ("each", "m", "mm", "m2"):
        problems.append("unit_unsupported")
    if values["currency"] != "AUD":
        problems.append("currency_unsupported")
    if values["tax_basis"] not in ("excluded", "GST Exclusive"):
        problems.append("tax_basis_unsupported")
    return values, list(dict.fromkeys(problems))


def preview_rows(
    document: dict[str, Any], sheet_index: int, header_row: int, mapping: dict[str, int | None]
) -> list[dict[str, Any]]:
    if type(sheet_index) is not int or not 1 <= sheet_index <= len(document["sheets"]):
        raise ValueError("sheet")
    sheet = document["sheets"][sheet_index - 1]
    if type(header_row) is not int or not 1 <= header_row < sheet["rows"]:
        raise ValueError("header row")
    validate_mapping(mapping, sheet["columns"])
    cells = {(cell["row"], cell["column"]): cell for cell in sheet["cells"]}
    result = []
    for number in range(header_row + 1, sheet["rows"] + 1):
        fields = {
            key: cells.get((number, column)) if column is not None else None
            for key, column in mapping.items()
        }
        if not any(fields.values()):
            continue
        values, problems = row_values(fields)
        row = {
            "sheet": sheet["name"],
            "sheet_index": sheet_index,
            "row": number,
            "header_row": header_row,
            "mapping": mapping,
            "fields": fields,
            "values": values,
            "problems": problems,
        }
        row["sha256"] = digest(row)
        result.append(row)
    return result


def validate_dataset_kind(value: Any) -> str:
    if type(value) is not str or value not in DATASET_KINDS:
        raise ValueError("dataset kind")
    return value


def validate_price_meaning(value: Any) -> str:
    if type(value) is not str or value not in PRICE_MEANINGS:
        raise ValueError("price meaning")
    return value


def profile_definition(
    document: dict[str, Any],
    *,
    draft_scope_id: str,
    source_id: str,
    dataset_id: str,
    dataset_kind: str,
    dataset_version: int,
    source_sha256: str,
    source_size_bytes: int,
    original_filename: str,
    document_sha256: str,
    sheet_index: int,
    header_row: int,
    mapping: dict[str, int | None],
    price_meaning: str,
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    validate_dataset_kind(dataset_kind)
    validate_price_meaning(price_meaning)
    rows = preview_rows(document, sheet_index, header_row, mapping)
    sheet = document["sheets"][sheet_index - 1]
    cells = {(cell["row"], cell["column"]): cell for cell in sheet["cells"]}
    header_cells = {
        field: cells.get((header_row, column)) if column is not None else None
        for field, column in mapping.items()
    }
    problem_counts: dict[str, int] = {}
    for row in rows:
        for problem in row["problems"]:
            problem_counts[problem] = problem_counts.get(problem, 0) + 1
    gaps = [field + "_unmapped" for field in FIELDS if mapping[field] is None]
    if price_meaning == "unknown":
        gaps.append("price_meaning_unknown")
    if dataset_kind == "general_pricelist":
        gaps.append("item_kind_not_supported_by_current_mapping")
    else:
        gaps.extend(
            (
                "system_manufacturer_not_supported_by_current_mapping",
                "system_configuration_not_supported_by_current_mapping",
            )
        )
    if problem_counts:
        gaps.append("row_anomalies_require_review")
    definition = {
        "schema_version": PROFILE_SCHEMA,
        "draft_scope_id": draft_scope_id,
        "dataset": {
            "id": dataset_id,
            "kind": dataset_kind,
            "version": dataset_version,
        },
        "source": {
            "id": source_id,
            "sha256": source_sha256,
            "size_bytes": source_size_bytes,
            "original_filename": original_filename,
            "document_sha256": document_sha256,
        },
        "selection": {
            "sheet_index": sheet_index,
            "sheet_name": sheet["name"],
            "header_row": header_row,
            "mapping": mapping,
            "header_cells": header_cells,
        },
        "commercial_basis": {"price_meaning": price_meaning},
        "diagnostics": {
            "sheet_rows": sheet["rows"],
            "sheet_columns": sheet["columns"],
            "data_rows": len(rows),
            "usable_rate_rows": sum(not row["problems"] for row in rows),
            "unresolved_rows": sum(bool(row["problems"]) for row in rows),
            "problem_counts": dict(sorted(problem_counts.items())),
            "mapped_fields": [field for field in FIELDS if mapping[field] is not None],
            "unmapped_fields": [field for field in FIELDS if mapping[field] is None],
            "gaps": list(dict.fromkeys(gaps)),
        },
        "approval_status": "unapproved",
        "effects": {
            "library_activated": False,
            "estimate_changed": False,
            "system_matching_performed": False,
            "price_inference_performed": False,
        },
    }
    validate_profile_definition(definition)
    return definition, rows


def _profile_id(value: Any) -> None:
    if type(value) is not str or not 1 <= len(value) <= 36:
        raise ValueError("profile identity")


def _profile_hash(value: Any, *, optional: bool = False) -> None:
    if optional and value is None:
        return
    if type(value) is not str or re.fullmatch(r"[0-9a-f]{64}", value) is None:
        raise ValueError("profile hash")


def _profile_cell(value: Any, row: int, column: int | None) -> None:
    if column is None:
        if value is not None:
            raise ValueError("profile header cell")
        return
    if type(value) is not dict or set(value) != {"address", "row", "column", "kind", "value"}:
        raise ValueError("profile header cell")
    from openpyxl.utils.cell import get_column_letter  # type: ignore[import-untyped]

    if (
        value["row"] != row
        or value["column"] != column
        or value["address"] != get_column_letter(column) + str(row)
        or value["kind"] not in ("text", "number", "date", "formula", "boolean", "error")
        or type(value["value"]) is not str
        or len(value["value"]) > 4000
    ):
        raise ValueError("profile header cell")


def validate_profile_definition(value: Any) -> None:
    if type(value) is not dict or set(value) != {
        "schema_version",
        "draft_scope_id",
        "dataset",
        "source",
        "selection",
        "commercial_basis",
        "diagnostics",
        "approval_status",
        "effects",
    }:
        raise ValueError("profile definition")
    if value["schema_version"] != PROFILE_SCHEMA or value["approval_status"] != "unapproved":
        raise ValueError("profile authority")
    _profile_id(value["draft_scope_id"])
    dataset = value["dataset"]
    if type(dataset) is not dict or set(dataset) != {"id", "kind", "version"}:
        raise ValueError("profile dataset")
    _profile_id(dataset["id"])
    validate_dataset_kind(dataset["kind"])
    if type(dataset["version"]) is not int or dataset["version"] < 1:
        raise ValueError("profile dataset version")
    source = value["source"]
    if type(source) is not dict or set(source) != {
        "id",
        "sha256",
        "size_bytes",
        "original_filename",
        "document_sha256",
    }:
        raise ValueError("profile source")
    _profile_id(source["id"])
    _profile_hash(source["sha256"])
    _profile_hash(source["document_sha256"])
    if type(source["size_bytes"]) is not int or not 1 <= source["size_bytes"] <= 10485760:
        raise ValueError("profile source size")
    if (
        type(source["original_filename"]) is not str
        or not 1 <= len(source["original_filename"]) <= 200
    ):
        raise ValueError("profile source name")
    selection = value["selection"]
    if type(selection) is not dict or set(selection) != {
        "sheet_index",
        "sheet_name",
        "header_row",
        "mapping",
        "header_cells",
    }:
        raise ValueError("profile selection")
    if type(selection["sheet_index"]) is not int or not 1 <= selection["sheet_index"] <= 10:
        raise ValueError("profile sheet")
    if type(selection["sheet_name"]) is not str or not 1 <= len(selection["sheet_name"]) <= 31:
        raise ValueError("profile sheet")
    if type(selection["header_row"]) is not int or not 1 <= selection["header_row"] <= 1000:
        raise ValueError("profile header")
    validate_mapping(selection["mapping"], 50)
    if type(selection["header_cells"]) is not dict or set(selection["header_cells"]) != set(FIELDS):
        raise ValueError("profile header cells")
    for field in FIELDS:
        _profile_cell(
            selection["header_cells"][field],
            selection["header_row"],
            selection["mapping"][field],
        )
    basis = value["commercial_basis"]
    if type(basis) is not dict or set(basis) != {"price_meaning"}:
        raise ValueError("profile basis")
    validate_price_meaning(basis["price_meaning"])
    diagnostics = value["diagnostics"]
    if type(diagnostics) is not dict or set(diagnostics) != {
        "sheet_rows",
        "sheet_columns",
        "data_rows",
        "usable_rate_rows",
        "unresolved_rows",
        "problem_counts",
        "mapped_fields",
        "unmapped_fields",
        "gaps",
    }:
        raise ValueError("profile diagnostics")
    for key, maximum in (("sheet_rows", 1000), ("sheet_columns", 50), ("data_rows", 999)):
        if type(diagnostics[key]) is not int or not 0 <= diagnostics[key] <= maximum:
            raise ValueError("profile diagnostics count")
    for key in ("usable_rate_rows", "unresolved_rows"):
        if type(diagnostics[key]) is not int or diagnostics[key] < 0:
            raise ValueError("profile diagnostics count")
    if diagnostics["usable_rate_rows"] + diagnostics["unresolved_rows"] != diagnostics["data_rows"]:
        raise ValueError("profile diagnostics total")
    if (
        type(diagnostics["mapped_fields"]) is not list
        or diagnostics["mapped_fields"]
        != [field for field in FIELDS if selection["mapping"][field]]
        or type(diagnostics["unmapped_fields"]) is not list
        or diagnostics["unmapped_fields"]
        != [field for field in FIELDS if not selection["mapping"][field]]
        or type(diagnostics["problem_counts"]) is not dict
        or any(
            type(key) is not str or type(count) is not int or count < 1
            for key, count in diagnostics["problem_counts"].items()
        )
        or type(diagnostics["gaps"]) is not list
        or len(diagnostics["gaps"]) != len(set(diagnostics["gaps"]))
        or any(type(gap) is not str or not 1 <= len(gap) <= 100 for gap in diagnostics["gaps"])
    ):
        raise ValueError("profile diagnostics")
    if value["effects"] != {
        "library_activated": False,
        "estimate_changed": False,
        "system_matching_performed": False,
        "price_inference_performed": False,
    }:
        raise ValueError("profile effects")
    if len(canonical(value)) > 131072:
        raise ValueError("profile size")


def validate_profile_envelope(value: Any) -> None:
    if type(value) is not dict or set(value) != {
        "schema_version",
        "profile_id",
        "revision",
        "parent_profile_sha256",
        "created_at",
        "created_by_id",
        "definition",
        "definition_sha256",
    }:
        raise ValueError("profile envelope")
    if value["schema_version"] != PROFILE_SCHEMA:
        raise ValueError("profile schema")
    _profile_id(value["profile_id"])
    _profile_id(value["created_by_id"])
    if type(value["revision"]) is not int or value["revision"] < 1:
        raise ValueError("profile revision")
    _profile_hash(value["parent_profile_sha256"], optional=True)
    if type(value["created_at"]) is not str or not 1 <= len(value["created_at"]) <= 64:
        raise ValueError("profile timestamp")
    validate_profile_definition(value["definition"])
    _profile_hash(value["definition_sha256"])
    if value["definition_sha256"] != digest(value["definition"]):
        raise ValueError("profile definition hash")
    if len(canonical(value)) > 131072:
        raise ValueError("profile size")


def validate_profile_decision_envelope(value: Any) -> None:
    if type(value) is not dict or set(value) != {
        "schema_version",
        "decision_id",
        "draft_scope_id",
        "source_id",
        "profile_id",
        "profile_revision",
        "profile_sha256",
        "decision",
        "reason",
        "reviewed_at",
        "reviewed_by_id",
        "effects",
    }:
        raise ValueError("profile decision envelope")
    if value["schema_version"] != PROFILE_DECISION_SCHEMA:
        raise ValueError("profile decision schema")
    for key in ("decision_id", "draft_scope_id", "source_id", "profile_id", "reviewed_by_id"):
        _profile_id(value[key])
    if type(value["profile_revision"]) is not int or value["profile_revision"] < 1:
        raise ValueError("profile decision revision")
    _profile_hash(value["profile_sha256"])
    if value["decision"] not in PROFILE_DECISIONS:
        raise ValueError("profile decision")
    reason = value["reason"]
    if type(reason) is not str or reason != reason.strip() or not 1 <= len(reason) <= 4000:
        raise ValueError("profile decision reason")
    try:
        reviewed_at = datetime.fromisoformat(value["reviewed_at"])
    except (TypeError, ValueError) as exc:
        raise ValueError("profile decision timestamp") from exc
    if reviewed_at.tzinfo is None or reviewed_at.utcoffset() != UTC.utcoffset(reviewed_at):
        raise ValueError("profile decision timestamp")
    if value["effects"] != {
        "rows_ingested": False,
        "library_activated": False,
        "technical_approval_granted": False,
        "system_matching_performed": False,
        "price_inference_performed": False,
        "estimate_changed": False,
        "release_performed": False,
    }:
        raise ValueError("profile decision effects")
    if len(canonical(value)) > 16384:
        raise ValueError("profile decision size")


def _validate_observation_row(row: Any) -> None:
    if type(row) is not dict or set(row) != {
        "sheet",
        "sheet_index",
        "row",
        "header_row",
        "mapping",
        "fields",
        "values",
        "problems",
        "sha256",
    }:
        raise ValueError("observation row")
    if type(row["sheet"]) is not str or not 1 <= len(row["sheet"]) <= 31:
        raise ValueError("observation sheet")
    for key, maximum in (("sheet_index", 10), ("row", 1000), ("header_row", 1000)):
        if type(row[key]) is not int or not 1 <= row[key] <= maximum:
            raise ValueError("observation row position")
    if row["row"] <= row["header_row"]:
        raise ValueError("observation row position")
    validate_mapping(row["mapping"], 50)
    if type(row["fields"]) is not dict or set(row["fields"]) != set(FIELDS):
        raise ValueError("observation fields")
    for key, cell in row["fields"].items():
        if cell is None:
            continue
        if type(cell) is not dict or set(cell) != {"address", "row", "column", "kind", "value"}:
            raise ValueError("observation cell")
        from openpyxl.utils.cell import get_column_letter  # type: ignore[import-untyped]

        if (
            type(cell["row"]) is not int
            or cell["row"] != row["row"]
            or type(cell["column"]) is not int
            or cell["column"] != row["mapping"][key]
            or cell["address"] != get_column_letter(cell["column"]) + str(cell["row"])
            or cell["kind"] not in ("text", "number", "date", "boolean")
            or type(cell["value"]) is not str
            or len(cell["value"]) > 4000
        ):
            raise ValueError("observation cell")
    values, problems = row_values(row["fields"])
    if problems or row["values"] != values or row["problems"] != []:
        raise ValueError("unusable observation row")
    if row["sha256"] != digest({key: item for key, item in row.items() if key != "sha256"}):
        raise ValueError("observation row hash")


def row_observation_definition(
    *,
    draft_scope_id: str,
    dataset: dict[str, Any],
    source: dict[str, Any],
    profile: dict[str, Any],
    row: dict[str, Any],
    item_kind: str,
    normalized_reference: str,
    evidence_state: str,
    review_reason: str,
    unresolved_fields: list[str],
) -> dict[str, Any]:
    value = {
        "schema_version": ROW_OBSERVATION_SCHEMA,
        "draft_scope_id": draft_scope_id,
        "dataset": dataset,
        "source": source,
        "profile": profile,
        "row": row,
        "interpretation": {
            "item_kind": item_kind,
            "normalized_reference": (
                normalized_reference.strip()
                if type(normalized_reference) is str
                else normalized_reference
            ),
            "evidence_state": evidence_state,
            "review_reason": review_reason.strip() if type(review_reason) is str else review_reason,
            "unresolved_fields": unresolved_fields,
        },
        "effects": {
            "rows_ingested": False,
            "product_created": False,
            "material_created": False,
            "labour_created": False,
            "service_created": False,
            "library_activated": False,
            "technical_approval_granted": False,
            "system_matching_performed": False,
            "price_inference_performed": False,
            "estimate_changed": False,
            "release_performed": False,
        },
    }
    validate_row_observation_definition(value)
    return value


def validate_row_observation_definition(value: Any) -> None:
    if type(value) is not dict or set(value) != {
        "schema_version",
        "draft_scope_id",
        "dataset",
        "source",
        "profile",
        "row",
        "interpretation",
        "effects",
    }:
        raise ValueError("row observation definition")
    if value["schema_version"] != ROW_OBSERVATION_SCHEMA:
        raise ValueError("row observation schema")
    _profile_id(value["draft_scope_id"])
    dataset = value["dataset"]
    if type(dataset) is not dict or set(dataset) != {"id", "kind", "version"}:
        raise ValueError("row observation dataset")
    _profile_id(dataset["id"])
    if dataset["kind"] != "general_pricelist":
        raise ValueError("row observation dataset kind")
    if type(dataset["version"]) is not int or dataset["version"] < 1:
        raise ValueError("row observation dataset version")
    source = value["source"]
    if type(source) is not dict or set(source) != {
        "id",
        "sha256",
        "size_bytes",
        "document_sha256",
    }:
        raise ValueError("row observation source")
    _profile_id(source["id"])
    _profile_hash(source["sha256"])
    _profile_hash(source["document_sha256"])
    if type(source["size_bytes"]) is not int or not 1 <= source["size_bytes"] <= 10485760:
        raise ValueError("row observation source size")
    profile = value["profile"]
    if type(profile) is not dict or set(profile) != {
        "id",
        "revision",
        "sha256",
        "decision_id",
        "decision_sha256",
        "price_meaning",
    }:
        raise ValueError("row observation profile")
    _profile_id(profile["id"])
    _profile_id(profile["decision_id"])
    _profile_hash(profile["sha256"])
    _profile_hash(profile["decision_sha256"])
    if type(profile["revision"]) is not int or profile["revision"] < 1:
        raise ValueError("row observation profile revision")
    if profile["price_meaning"] == "unknown":
        raise ValueError("row observation price meaning")
    validate_price_meaning(profile["price_meaning"])
    _validate_observation_row(value["row"])
    interpretation = value["interpretation"]
    if type(interpretation) is not dict or set(interpretation) != {
        "item_kind",
        "normalized_reference",
        "evidence_state",
        "review_reason",
        "unresolved_fields",
    }:
        raise ValueError("row observation interpretation")
    if interpretation["item_kind"] not in ROW_ITEM_KINDS:
        raise ValueError("row observation item kind")
    if interpretation["evidence_state"] not in ROW_EVIDENCE_STATES:
        raise ValueError("row observation evidence state")
    for key, maximum in (("normalized_reference", 300), ("review_reason", 4000)):
        item = interpretation[key]
        if type(item) is not str or item != item.strip() or not 1 <= len(item) <= maximum:
            raise ValueError("row observation interpretation")
    unresolved = interpretation["unresolved_fields"]
    if (
        type(unresolved) is not list
        or unresolved != list(dict.fromkeys(unresolved))
        or any(type(item) is not str or item not in FIELDS for item in unresolved)
    ):
        raise ValueError("row observation unresolved fields")
    if interpretation["evidence_state"] == "confirmed" and unresolved:
        raise ValueError("confirmed row observation unresolved")
    if value["effects"] != {
        "rows_ingested": False,
        "product_created": False,
        "material_created": False,
        "labour_created": False,
        "service_created": False,
        "library_activated": False,
        "technical_approval_granted": False,
        "system_matching_performed": False,
        "price_inference_performed": False,
        "estimate_changed": False,
        "release_performed": False,
    }:
        raise ValueError("row observation effects")
    if len(canonical(value)) > 131072:
        raise ValueError("row observation size")


def validate_row_observation_envelope(value: Any) -> None:
    if type(value) is not dict or set(value) != {
        "schema_version",
        "observation_id",
        "reviewed_at",
        "reviewed_by_id",
        "definition",
        "definition_sha256",
    }:
        raise ValueError("row observation envelope")
    if value["schema_version"] != ROW_OBSERVATION_SCHEMA:
        raise ValueError("row observation schema")
    _profile_id(value["observation_id"])
    _profile_id(value["reviewed_by_id"])
    try:
        reviewed_at = datetime.fromisoformat(value["reviewed_at"])
    except (TypeError, ValueError) as exc:
        raise ValueError("row observation timestamp") from exc
    if reviewed_at.tzinfo is None or reviewed_at.utcoffset() != UTC.utcoffset(reviewed_at):
        raise ValueError("row observation timestamp")
    validate_row_observation_definition(value["definition"])
    _profile_hash(value["definition_sha256"])
    if value["definition_sha256"] != digest(value["definition"]):
        raise ValueError("row observation definition hash")
    if len(canonical(value)) > 131072:
        raise ValueError("row observation size")


def _system_mapping_variant_snapshot(value: Any) -> None:
    keys = {
        "id",
        "key",
        "variant_id",
        "system_id",
        "record_version",
        "source_hash",
        "technical_fields",
        "technical_fields_sha256",
        "release_record",
        "release_record_sha256",
        "source_verification",
        "sha256",
    }
    if type(value) is not dict or set(value) != keys:
        raise ValueError("system mapping variant")
    _profile_id(value["id"])
    for key in ("key", "variant_id", "system_id"):
        item = value[key]
        if type(item) is not str or item != item.strip() or not 1 <= len(item) <= 300:
            raise ValueError("system mapping variant identity")
    if type(value["record_version"]) is not int or value["record_version"] < 1:
        raise ValueError("system mapping variant revision")
    _profile_hash(value["source_hash"])
    fields = value["technical_fields"]
    if (
        type(fields) is not dict
        or set(fields) != set(TECHNICAL_FIELD_NAMES)
        or any(
            item is not None
            and type(item) not in (str, bool)
            or type(item) is str
            and len(item) > 4000
            for item in fields.values()
        )
    ):
        raise ValueError("system mapping technical fields")
    _profile_hash(value["technical_fields_sha256"])
    if value["technical_fields_sha256"] != digest(fields):
        raise ValueError("system mapping technical fields hash")
    record = value["release_record"]
    if type(record) is not dict or set(record) != {
        "id",
        "key",
        "variant_id",
        "system_id",
        "frl",
        "source_document_reference",
        "source_page",
        "source_hash",
        "record_version",
        "technical_fields",
        "source_binding",
    }:
        raise ValueError("system mapping release record")
    if (
        record["id"] != value["id"]
        or record["key"] != value["key"]
        or record["variant_id"] != value["variant_id"]
        or record["system_id"] != value["system_id"]
        or record["source_hash"] != value["source_hash"]
        or record["record_version"] != value["record_version"]
        or record["technical_fields"] != fields
    ):
        raise ValueError("system mapping release binding")
    for key, maximum in (
        ("frl", 100),
        ("source_document_reference", 300),
        ("source_page", 100),
    ):
        item = record[key]
        if item is not None and (type(item) is not str or len(item) > maximum):
            raise ValueError("system mapping release record")
    validate_binding(record["source_binding"])
    if record["source_binding"]["state"] != "bound":
        raise ValueError("system mapping source binding")
    _profile_hash(value["release_record_sha256"])
    if value["release_record_sha256"] != digest(record):
        raise ValueError("system mapping release record hash")
    verification = value["source_verification"]
    if type(verification) is not dict or set(verification) != {"method"}:
        raise ValueError("system mapping source verification")
    if verification["method"] != "exact_bytes":
        raise ValueError("system mapping source verification")
    _profile_hash(value["sha256"])
    if value["sha256"] != digest({key: item for key, item in value.items() if key != "sha256"}):
        raise ValueError("system mapping variant hash")


def system_mapping_definition(
    *,
    draft_scope_id: str,
    dataset: dict[str, Any],
    source: dict[str, Any],
    profile: dict[str, Any],
    row: dict[str, Any],
    technical_release: dict[str, Any],
    variants: list[dict[str, Any]],
    mapping_status: str,
    normalized_reference: str,
    selected_variant_id: str | None,
    candidate_variant_ids: list[str],
    identity_evidence_fields: list[str],
    review_reason: str,
    unresolved_fields: list[str],
) -> dict[str, Any]:
    value = {
        "schema_version": SYSTEM_MAPPING_SCHEMA,
        "draft_scope_id": draft_scope_id,
        "dataset": dataset,
        "source": source,
        "profile": profile,
        "row": row,
        "technical_release": technical_release,
        "variants": variants,
        "interpretation": {
            "status": mapping_status,
            "normalized_reference": (
                normalized_reference.strip()
                if type(normalized_reference) is str
                else normalized_reference
            ),
            "selected_variant_id": selected_variant_id,
            "candidate_variant_ids": candidate_variant_ids,
            "identity_evidence_fields": identity_evidence_fields,
            "review_reason": review_reason.strip() if type(review_reason) is str else review_reason,
            "unresolved_fields": unresolved_fields,
        },
        "effects": {
            "pricing_system_mapping_recorded": True,
            "rows_ingested": False,
            "library_activated": False,
            "technical_approval_granted": False,
            "technical_applicability_assessed": False,
            "system_match_package_changed": False,
            "price_inference_performed": False,
            "estimate_changed": False,
            "holdout_assigned": False,
            "release_performed": False,
        },
    }
    validate_system_mapping_definition(value)
    return value


def validate_system_mapping_definition(value: Any) -> None:
    if type(value) is not dict or set(value) != {
        "schema_version",
        "draft_scope_id",
        "dataset",
        "source",
        "profile",
        "row",
        "technical_release",
        "variants",
        "interpretation",
        "effects",
    }:
        raise ValueError("system mapping definition")
    if value["schema_version"] != SYSTEM_MAPPING_SCHEMA:
        raise ValueError("system mapping schema")
    _profile_id(value["draft_scope_id"])
    dataset = value["dataset"]
    if (
        type(dataset) is not dict
        or set(dataset) != {"id", "kind", "version"}
        or dataset["kind"] != "firefly_system_prices"
    ):
        raise ValueError("system mapping dataset")
    _profile_id(dataset["id"])
    if type(dataset["version"]) is not int or dataset["version"] < 1:
        raise ValueError("system mapping dataset version")
    source = value["source"]
    if type(source) is not dict or set(source) != {
        "id",
        "sha256",
        "size_bytes",
        "document_sha256",
    }:
        raise ValueError("system mapping source")
    _profile_id(source["id"])
    _profile_hash(source["sha256"])
    _profile_hash(source["document_sha256"])
    if type(source["size_bytes"]) is not int or not 1 <= source["size_bytes"] <= 10485760:
        raise ValueError("system mapping source size")
    profile = value["profile"]
    if type(profile) is not dict or set(profile) != {
        "id",
        "revision",
        "sha256",
        "decision_id",
        "decision_sha256",
        "price_meaning",
    }:
        raise ValueError("system mapping profile")
    _profile_id(profile["id"])
    _profile_id(profile["decision_id"])
    _profile_hash(profile["sha256"])
    _profile_hash(profile["decision_sha256"])
    if type(profile["revision"]) is not int or profile["revision"] < 1:
        raise ValueError("system mapping profile revision")
    if profile["price_meaning"] == "unknown":
        raise ValueError("system mapping price meaning")
    validate_price_meaning(profile["price_meaning"])
    _validate_observation_row(value["row"])
    release = value["technical_release"]
    if type(release) is not dict or set(release) != {
        "id",
        "version",
        "sha256",
        "effective_date",
    }:
        raise ValueError("system mapping technical release")
    _profile_id(release["id"])
    _profile_hash(release["sha256"])
    if (
        type(release["version"]) is not str
        or release["version"] != release["version"].strip()
        or not 1 <= len(release["version"]) <= 50
    ):
        raise ValueError("system mapping technical release")
    if release["effective_date"] is not None:
        try:
            date.fromisoformat(release["effective_date"])
        except (TypeError, ValueError) as exc:
            raise ValueError("system mapping technical release") from exc
    variants = value["variants"]
    if type(variants) is not list or len(variants) > 20:
        raise ValueError("system mapping variants")
    for variant in variants:
        _system_mapping_variant_snapshot(variant)
    variant_ids = [variant["id"] for variant in variants]
    if variant_ids != sorted(set(variant_ids)):
        raise ValueError("system mapping variants")
    interpretation = value["interpretation"]
    if type(interpretation) is not dict or set(interpretation) != {
        "status",
        "normalized_reference",
        "selected_variant_id",
        "candidate_variant_ids",
        "identity_evidence_fields",
        "review_reason",
        "unresolved_fields",
    }:
        raise ValueError("system mapping interpretation")
    status = interpretation["status"]
    if status not in SYSTEM_MAPPING_STATES:
        raise ValueError("system mapping status")
    for key, maximum in (("normalized_reference", 300), ("review_reason", 4000)):
        item = interpretation[key]
        if type(item) is not str or item != item.strip() or not 1 <= len(item) <= maximum:
            raise ValueError("system mapping interpretation")
    selected = interpretation["selected_variant_id"]
    if selected is not None:
        _profile_id(selected)
    candidates = interpretation["candidate_variant_ids"]
    if (
        type(candidates) is not list
        or candidates != sorted(set(candidates))
        or len(candidates) > 20
        or any(type(item) is not str or not 1 <= len(item) <= 36 for item in candidates)
        or candidates != variant_ids
    ):
        raise ValueError("system mapping candidates")
    evidence = interpretation["identity_evidence_fields"]
    if (
        type(evidence) is not list
        or evidence != list(dict.fromkeys(evidence))
        or any(item not in SYSTEM_IDENTITY_EVIDENCE_FIELDS for item in evidence)
        or any(not value["row"]["values"].get(item) for item in evidence)
    ):
        raise ValueError("system mapping identity evidence")
    unresolved = interpretation["unresolved_fields"]
    if (
        type(unresolved) is not list
        or unresolved != list(dict.fromkeys(unresolved))
        or any(item not in SYSTEM_MAPPING_UNRESOLVED_FIELDS for item in unresolved)
    ):
        raise ValueError("system mapping unresolved fields")
    if status == "mapped":
        if selected is None or candidates != [selected] or len(variants) != 1 or not evidence:
            raise ValueError("mapped system identity")
        if unresolved:
            raise ValueError("mapped system unresolved")
    elif status == "ambiguous":
        if selected is not None or len(candidates) < 2 or not evidence or not unresolved:
            raise ValueError("ambiguous system identity")
    elif selected is not None or candidates or variants or not unresolved:
        raise ValueError("unmatched system identity")
    if value["effects"] != {
        "pricing_system_mapping_recorded": True,
        "rows_ingested": False,
        "library_activated": False,
        "technical_approval_granted": False,
        "technical_applicability_assessed": False,
        "system_match_package_changed": False,
        "price_inference_performed": False,
        "estimate_changed": False,
        "holdout_assigned": False,
        "release_performed": False,
    }:
        raise ValueError("system mapping effects")
    if len(canonical(value)) > 524288:
        raise ValueError("system mapping size")


def validate_system_mapping_envelope(value: Any) -> None:
    if type(value) is not dict or set(value) != {
        "schema_version",
        "mapping_id",
        "reviewed_at",
        "reviewed_by_id",
        "definition",
        "definition_sha256",
    }:
        raise ValueError("system mapping envelope")
    if value["schema_version"] != SYSTEM_MAPPING_SCHEMA:
        raise ValueError("system mapping schema")
    _profile_id(value["mapping_id"])
    _profile_id(value["reviewed_by_id"])
    try:
        reviewed_at = datetime.fromisoformat(value["reviewed_at"])
    except (TypeError, ValueError) as exc:
        raise ValueError("system mapping timestamp") from exc
    if reviewed_at.tzinfo is None or reviewed_at.utcoffset() != UTC.utcoffset(reviewed_at):
        raise ValueError("system mapping timestamp")
    validate_system_mapping_definition(value["definition"])
    _profile_hash(value["definition_sha256"])
    if value["definition_sha256"] != digest(value["definition"]):
        raise ValueError("system mapping definition hash")
    if len(canonical(value)) > 524288:
        raise ValueError("system mapping size")


def validate_selection(selection: Any, lines: list[dict[str, Any]]) -> None:
    if type(selection) is not dict or set(selection) != {
        "line_id",
        "event_id",
        "source_id",
        "source_sha256",
        "source_size_bytes",
        "original_filename",
        "document_sha256",
        "scan_sha256",
        "row",
        "method",
        "approval_status",
        "recovery_note",
    }:
        raise ValueError("rate provenance")
    if selection["method"] != "exact_library_rate" or selection["approval_status"] != "unreviewed":
        raise ValueError("rate authority")
    for key in ("source_sha256", "document_sha256", "scan_sha256"):
        if type(selection[key]) is not str or re.fullmatch(r"[0-9a-f]{64}", selection[key]) is None:
            raise ValueError("source hash")
    if (
        type(selection["source_size_bytes"]) is not int
        or not 1 <= selection["source_size_bytes"] <= 10485760
    ):
        raise ValueError("source size")
    for key, bound in (("source_id", 36), ("original_filename", 200), ("recovery_note", 4000)):
        if type(selection[key]) is not str or not 1 <= len(selection[key].strip()) <= bound:
            raise ValueError("source metadata")
    row = selection["row"]
    if type(row) is not dict or set(row) != {
        "sheet",
        "sheet_index",
        "row",
        "header_row",
        "mapping",
        "fields",
        "values",
        "problems",
        "sha256",
    }:
        raise ValueError("row")
    if type(row["sheet"]) is not str or not 1 <= len(row["sheet"]) <= 31:
        raise ValueError("sheet name")
    for key, maximum in (("sheet_index", 10), ("row", 1000), ("header_row", 1000)):
        if type(row[key]) is not int or not 1 <= row[key] <= maximum:
            raise ValueError("row position")
    if row["row"] <= row["header_row"]:
        raise ValueError("row position")
    validate_mapping(row["mapping"], 50)
    if type(row["fields"]) is not dict or set(row["fields"]) != set(FIELDS):
        raise ValueError("source fields")
    for key, cell in row["fields"].items():
        if cell is None:
            continue
        if type(cell) is not dict or set(cell) != {"address", "row", "column", "kind", "value"}:
            raise ValueError("source cell")
        if (
            type(cell["row"]) is not int
            or cell["row"] != row["row"]
            or type(cell["column"]) is not int
            or cell["column"] != row["mapping"][key]
            or type(cell["address"]) is not str
            or re.fullmatch(r"[A-Z]{1,2}[1-9][0-9]{0,3}", cell["address"]) is None
            or cell["kind"] not in ("text", "number", "date", "formula", "boolean", "error")
            or type(cell["value"]) is not str
            or len(cell["value"]) > 4000
        ):
            raise ValueError("source cell")
        from openpyxl.utils.cell import get_column_letter  # type: ignore[import-untyped]

        if cell["address"] != get_column_letter(cell["column"]) + str(cell["row"]):
            raise ValueError("cell address")
    values, problems = row_values(row["fields"])
    if problems or row["values"] != values or row["problems"] != []:
        raise ValueError("unusable rate")
    if row["sha256"] != digest({key: value for key, value in row.items() if key != "sha256"}):
        raise ValueError("row hash")
    line = next((item for item in lines if item["line_id"] == selection["line_id"]), None)
    event = (
        next((item for item in line["history"] if item["event_id"] == selection["event_id"]), None)
        if line
        else None
    )
    if (
        line is None
        or event is None
        or event["action"] != "override"
        or event["unit_sell_rate"] != values["rate"]
        or line["unit"] != values["unit"]
        or event["reason"] != selection["recovery_note"]
    ):
        raise ValueError("rate event")
    if len(canonical(selection)) > 65536:
        raise ValueError("selection size")
