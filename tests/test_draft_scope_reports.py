from __future__ import annotations

import copy
import hashlib
import io
import json

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import create_engine, func, inspect, select, update
from sqlalchemy.orm import sessionmaker
from test_draft_scope import sample_payload, uid

from classifire import physical_models  # noqa: F401
from classifire.db import Base
from classifire.models import (
    AuditEvent,
    DraftScopeReport,
    DraftScopeRevision,
    Estimate,
    Opening,
    Project,
    Service,
    User,
)
from classifire.outputs import draft_scope as renderers
from classifire.physical_models import PhysicalModelLock
from classifire.services.draft_scope import (
    apply_import,
    create_draft_project,
    read_revision,
    revision_bytes,
    save_revision,
)
from classifire.services.draft_scope_reports import (
    MAX_OUTPUT_BYTES,
    REPORT_LIST_LIMIT,
    REPORT_SCHEMA_VERSION,
    DraftScopeReportError,
    create_report,
    list_reports,
    read_report,
    report_bytes,
    report_freshness,
    validate_report_snapshot,
)


@pytest.fixture
def case(tmp_path):
    engine = create_engine("sqlite:///" + str(tmp_path / "reports.sqlite"))
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    with factory() as db:
        db.add_all(
            [
                User(
                    id=uid(number),
                    email=f"report-{number}@example.test",
                    full_name=f"User {number}",
                    password_hash="unused",  # noqa: S106 - synthetic non-authenticating fixture
                    role=role,
                    is_active=True,
                )  # noqa: S106
                for number, role in (
                    (100, "estimator"),
                    (101, "estimator"),
                    (102, "administrator"),
                    (103, "read_only"),
                )
            ]
        )
        db.commit()
        actor = db.get(User, uid(100))
        draft = create_draft_project(db, actor, "SYNTHETIC-REPORT", "Synthetic report project")
        save_revision(db, actor, draft.id, 1, sample_payload())
        draft_id = draft.id
        db.commit()
    yield factory, draft_id
    engine.dispose()


def _counts(db):
    return {
        model.__tablename__: db.scalar(select(func.count()).select_from(model))
        for model in (
            AuditEvent,
            DraftScopeReport,
            DraftScopeRevision,
            Estimate,
            Opening,
            Service,
            PhysicalModelLock,
        )
    }


def _fast_renderers(monkeypatch):
    monkeypatch.setattr(
        renderers,
        "render_scope_report_pdf",
        lambda snapshot: b"%PDF-synthetic" + snapshot["sha256"].encode(),
    )
    monkeypatch.setattr(
        renderers,
        "render_scope_report_xlsx",
        lambda snapshot: b"PK\x03\x04synthetic" + snapshot["sha256"].encode(),
    )


def test_real_pair_uses_one_snapshot_persists_restart_and_never_recalculates(case, tmp_path):
    factory, draft_id = case
    with factory() as db:
        actor = db.get(User, uid(100))
        before = _counts(db)
        source = read_revision(db, actor, draft_id, 2)
        report = create_report(db, actor, draft_id, 2)
        snapshot = read_report(db, actor, draft_id, report.id)
        validate_report_snapshot(snapshot)
        assert snapshot["scope"] == source
        assert snapshot["project"] == {
            "id": source["project_id"],
            "reference": "SYNTHETIC-REPORT",
            "name": "Synthetic report project",
        }
        assert snapshot["schema_version"] == REPORT_SCHEMA_VERSION
        assert snapshot["profile"] == "scope-only"
        assert snapshot["state"] == "Draft" and snapshot["review_status"] == "unreviewed"
        pdf = report_bytes(db, actor, draft_id, report.id, "pdf")
        xlsx = report_bytes(db, actor, draft_id, report.id, "xlsx")
        (tmp_path / "synthetic-scope-report.pdf").write_bytes(pdf)
        (tmp_path / "synthetic-scope-report.xlsx").write_bytes(xlsx)
        assert len(PdfReader(io.BytesIO(pdf)).pages) >= 1
        workbook = load_workbook(io.BytesIO(xlsx))
        assert len(workbook.sheetnames) >= 1
        workbook.close()
        assert hashlib.sha256(pdf).hexdigest() == report.pdf_sha256
        assert hashlib.sha256(xlsx).hexdigest() == report.xlsx_sha256
        assert report_freshness(db, actor, draft_id, report.id) is False
        after = _counts(db)
        for name in (
            "estimates",
            "openings",
            "services",
            "physical_model_locks",
            "draft_scope_revisions",
        ):
            assert after[name] == before[name]
        created_audit = db.scalar(
            select(AuditEvent).where(AuditEvent.action == "draft_scope_report.create")
        )
        assert set(created_audit.new_value) == {
            "scope_revision",
            "scope_sha256",
            "snapshot_sha256",
            "pdf_sha256",
            "xlsx_sha256",
        }
        report_id = report.id
        db.commit()
    with factory() as db:
        actor = db.get(User, uid(100))
        assert read_report(db, actor, draft_id, report_id) == snapshot
        assert report_bytes(db, actor, draft_id, report_id, "pdf") == pdf
        assert report_bytes(db, actor, draft_id, report_id, "xlsx") == xlsx


def test_old_outputs_survive_scope_and_project_edits_with_stale_status(case, monkeypatch):
    _fast_renderers(monkeypatch)
    factory, draft_id = case
    with factory() as db:
        actor = db.get(User, uid(100))
        first = create_report(db, actor, draft_id, 2)
        pdf, xlsx = first.pdf_bytes, first.xlsx_bytes
        snapshot = read_report(db, actor, draft_id, first.id)
        db.commit()
        project = db.get(Project, snapshot["project"]["id"])
        project.name = "Changed current name"
        db.commit()
        assert report_freshness(db, actor, draft_id, first.id) is True
        assert (
            read_report(db, actor, draft_id, first.id)["project"]["name"]
            == "Synthetic report project"
        )
        assert report_bytes(db, actor, draft_id, first.id, "pdf") == pdf
        project.name = "Synthetic report project"
        db.commit()
        assert report_freshness(db, actor, draft_id, first.id) is False
        save_revision(db, actor, draft_id, 2, {"assumptions": ["A newer scope"]})
        db.commit()
        assert report_freshness(db, actor, draft_id, first.id) is True
        assert report_bytes(db, actor, draft_id, first.id, "xlsx") == xlsx
        older = create_report(db, actor, draft_id, 1)
        assert older.scope_revision == 1
        assert report_freshness(db, actor, draft_id, older.id) is True


def test_imported_v2_provenance_is_frozen_in_report(case, monkeypatch):
    _fast_renderers(monkeypatch)
    factory, draft_id = case
    with factory() as db:
        actor = db.get(User, uid(100))
        raw = revision_bytes(db, actor, draft_id, 2)
        imported = apply_import(db, actor, draft_id, 2, raw, hashlib.sha256(raw).hexdigest())
        report = create_report(db, actor, draft_id, 3)
        assert read_report(db, actor, draft_id, report.id)["scope"] == imported
        assert imported["import_lineage"]
        assert imported["review_status"] == "unreviewed"


def test_report_permissions_and_object_binding(case, monkeypatch):
    _fast_renderers(monkeypatch)
    factory, draft_id = case
    with factory() as db:
        actor = db.get(User, uid(100))
        report = create_report(db, actor, draft_id, 2)
        db.commit()
        other = db.get(User, uid(101))
        for call in (
            lambda: create_report(db, other, draft_id, 2),
            lambda: list_reports(db, other, draft_id),
            lambda: read_report(db, other, draft_id, report.id),
            lambda: report_bytes(db, other, draft_id, report.id, "pdf"),
            lambda: report_freshness(db, other, draft_id, report.id),
        ):
            with pytest.raises(DraftScopeReportError, match="DRAFT_NOT_FOUND"):
                call()
        reader = db.get(User, uid(103))
        with pytest.raises(DraftScopeReportError, match="DRAFT_PERMISSION_DENIED"):
            create_report(db, reader, draft_id, 2)
        with pytest.raises(DraftScopeReportError, match="DRAFT_PERMISSION_DENIED"):
            create_report(db, object(), draft_id, 2)
        admin = db.get(User, uid(102))
        assert read_report(db, admin, draft_id, report.id)["report_id"] == report.id
        another = create_draft_project(db, actor, "ANOTHER", "Other target")
        with pytest.raises(DraftScopeReportError, match="DRAFT_REPORT_NOT_FOUND"):
            read_report(db, actor, another.id, report.id)
        actor.is_active = False
        db.commit()
        with pytest.raises(DraftScopeReportError, match="DRAFT_PERMISSION_DENIED"):
            report_bytes(db, actor, draft_id, report.id, "pdf")


@pytest.mark.parametrize("revision", [None, True, 0, -1, 999, "2"])
def test_report_requires_explicit_existing_saved_revision(case, monkeypatch, revision):
    _fast_renderers(monkeypatch)
    factory, draft_id = case
    with factory() as db:
        before = _counts(db)
        with pytest.raises(DraftScopeReportError, match="DRAFT_REVISION_NOT_FOUND"):
            create_report(db, db.get(User, uid(100)), draft_id, revision)
        assert _counts(db) == before


@pytest.mark.parametrize("failure", ["exception", "empty", "oversized", "mutated"])
def test_renderer_failure_retains_neither_output_nor_audit(case, monkeypatch, failure):
    _fast_renderers(monkeypatch)

    def fail(snapshot):
        if failure == "exception":
            raise RuntimeError("Private renderer details must not be echoed")
        if failure == "empty":
            return b""
        if failure == "oversized":
            return b"PK\x03\x04" + b"x" * MAX_OUTPUT_BYTES
        snapshot["project"]["name"] = "Mutated"
        return b"PK\x03\x04synthetic"

    monkeypatch.setattr(renderers, "render_scope_report_xlsx", fail)
    factory, draft_id = case
    with factory() as db:
        before = _counts(db)
        with pytest.raises(DraftScopeReportError) as raised:
            create_report(db, db.get(User, uid(100)), draft_id, 2)
        assert "Private" not in str(raised.value)
        assert _counts(db) == before


@pytest.mark.parametrize(
    "column,value",
    [
        ("snapshot_json", "{}"),
        ("snapshot_hash", "0" * 64),
        ("scope_hash", "0" * 64),
        ("pdf_bytes", b"%PDF-altered"),
        ("xlsx_bytes", b"PK\x03\x04altered"),
        ("pdf_sha256", "0" * 64),
        ("xlsx_sha256", "0" * 64),
        ("created_by_id", uid(101)),
        ("scope_revision", 1),
    ],
)
def test_report_integrity_failures_refuse_both_downloads(case, monkeypatch, column, value):
    _fast_renderers(monkeypatch)
    factory, draft_id = case
    with factory() as db:
        actor = db.get(User, uid(100))
        report = create_report(db, actor, draft_id, 2)
        db.execute(
            update(DraftScopeReport).where(DraftScopeReport.id == report.id).values({column: value})
        )
        db.commit()
        for format_name in ("pdf", "xlsx"):
            with pytest.raises(DraftScopeReportError, match="DRAFT_REPORT_INTEGRITY_FAILED"):
                report_bytes(db, actor, draft_id, report.id, format_name)


def test_report_checks_retained_source_integrity(case, monkeypatch):
    _fast_renderers(monkeypatch)
    factory, draft_id = case
    with factory() as db:
        actor = db.get(User, uid(100))
        report = create_report(db, actor, draft_id, 2)
        db.execute(
            update(DraftScopeRevision)
            .where(DraftScopeRevision.draft_scope_id == draft_id)
            .values(content_hash="0" * 64)
        )
        db.commit()
        with pytest.raises(DraftScopeReportError, match="DRAFT_REPORT_INTEGRITY_FAILED"):
            read_report(db, actor, draft_id, report.id)


def test_caller_rollback_discards_report_pair_and_create_audit(case, monkeypatch):
    _fast_renderers(monkeypatch)
    factory, draft_id = case
    with factory() as db:
        before = _counts(db)
        create_report(db, db.get(User, uid(100)), draft_id, 2)
        db.rollback()
        assert _counts(db) == before


def test_bounded_list_releases_binary_fields_and_keeps_old_reports_addressable(case, monkeypatch):
    _fast_renderers(monkeypatch)
    factory, draft_id = case
    with factory() as db:
        actor = db.get(User, uid(100))
        reports = [create_report(db, actor, draft_id, 2) for _ in range(REPORT_LIST_LIMIT + 1)]
        db.commit()
        recent = list_reports(db, actor, draft_id)
        assert len(recent) == REPORT_LIST_LIMIT
        assert reports[0].id not in {row.id for row in recent}
        assert all({"pdf_bytes", "xlsx_bytes"} <= inspect(row).unloaded for row in recent)
        assert read_report(db, actor, draft_id, reports[0].id)["report_id"] == reports[0].id


def test_snapshot_verifier_rejects_resealed_authority_and_source_changes(case, monkeypatch):
    _fast_renderers(monkeypatch)
    factory, draft_id = case
    with factory() as db:
        report = create_report(db, db.get(User, uid(100)), draft_id, 2)
        original = json.loads(report.snapshot_json)
    for field, value in (
        ("state", "Released"),
        ("review_status", "approved"),
        ("profile", "complete"),
        ("render_version", True),
    ):
        snapshot = copy.deepcopy(original)
        snapshot[field] = value
        snapshot.pop("sha256")
        snapshot["sha256"] = hashlib.sha256(
            json.dumps(snapshot, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode()
        ).hexdigest()
        with pytest.raises(DraftScopeReportError, match="DRAFT_REPORT_SNAPSHOT_INVALID"):
            validate_report_snapshot(snapshot)
