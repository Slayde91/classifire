from __future__ import annotations

import io
from copy import deepcopy

import pytest
from openpyxl import load_workbook
from test_draft_scope_outputs import _id, _pdf_text, _seal, _snapshot

from classifire.outputs import draft_estimate as estimate_outputs
from classifire.outputs import draft_scope as scope_outputs
from classifire.services import draft_estimate_contract as estimates
from classifire.services import draft_estimate_reports as estimate_reports
from classifire.services import draft_package_import as imports
from classifire.services import draft_project_packages as packages
from classifire.services import draft_scope_reports as scope_reports
from classifire.services import draft_system_match_contract as matches
from classifire.services.draft_scope import DraftScopeError
from classifire.services.draft_scope_evidence import observation_hash


def page_reference():
    return {
        "source_id": _id(800),
        "source_sha256": "a" * 64,
        "source_size_bytes": 900,
        "original_filename": "Synthetic <report>.pdf",
        "page_number": 1,
        "locator_key": "page-1",
        "page_text_sha256": "b" * 64,
        "document_sha256": "c" * 64,
        "scan_sha256": "d" * 64,
        "reviewed_by": _id(1),
        "reviewed_at": "2026-09-06T00:00:00+00:00",
        "method": "human_page_entity_review",
        "origin": "local_retained",
    }


def entity_scope():
    scope = _snapshot(imported=False)["scope"]
    refs = []
    for kind, item in (
        ("defect", scope["content"]["defects"][0]),
        ("opening", scope["content"]["openings"][0]),
        ("service", scope["content"]["services"][0]),
        ("service", scope["content"]["services"][1]),
    ):
        refs.append(
            page_reference()
            | {
                "target_kind": kind,
                "target_id": item["id"],
                "target_sha256": observation_hash(item),
            }
        )
    refs[2]["origin"] = "imported_unverified"
    observation = scope["content"]["observations"][0]
    refs.append(
        page_reference()
        | {
            "observation_id": observation["id"],
            "observation_sha256": observation_hash(observation),
            "method": "human_page_review",
        }
    )
    # These saved claims remain valid historical provenance after a later edit/removal.
    scope["content"]["openings"][0]["substrate"] = "Changed substrate"
    scope["content"]["services"].pop()
    scope.update(
        schema_version="CLASSIFIRE-DRAFT-SCOPE-v4",
        provenance="manual_edit",
        import_lineage=[],
        evidence_refs=refs,
    )
    return _seal(scope)


def review(scope):
    return _seal(
        {
            "schema_version": "CLASSIFIRE-DRAFT-SYSTEM-MATCH-v1",
            "artifact_id": _id(30),
            "project_id": scope["project_id"],
            "revision": 1,
            "parent_hash": None,
            "created_by": _id(1),
            "created_at": "2026-09-06T01:00:00+00:00",
            "state": "Draft",
            "review_status": "unreviewed",
            "provenance": "manual_review",
            "coverage": "selected_target_only",
            "scope": scope,
            "release": {
                "id": _id(31),
                "version": "synthetic",
                "sha256": "e" * 64,
                "effective_date": None,
            },
            "target": matches.target_for(scope, _id(5), None),
            "retrieval": {
                "version": 1,
                "inputs": dict.fromkeys(matches.COMPARISON_NAMES),
                "missing_criteria": [],
                "unassessed_criteria": ["complete_applicability"],
                "truncated": False,
                "candidate_limit": 20,
                "evaluated_at": "2026-09-06T01:00:00+00:00",
            },
            "candidates": [],
            "decisions": [],
        }
    )


def report_snapshot(profile, scope=None):
    scope = entity_scope() if scope is None else scope
    report = _snapshot(imported=False)
    report["scope"] = scope
    report["profile"] = profile
    if profile in ("scope-only", "scope-and-system"):
        report["render_version"] = 4 if profile == "scope-and-system" else 3
        if profile == "scope-and-system":
            report.update(
                schema_version=scope_reports.SYSTEM_REPORT_SCHEMA_VERSION,
                system_match=review(scope),
            )
    else:
        estimate = _seal(
            {
                "schema_version": estimates.SCHEMA_VERSION,
                "artifact_id": _id(40),
                "project_id": scope["project_id"],
                "revision": 1,
                "parent_hash": None,
                "created_by": _id(1),
                "created_at": "2026-09-06T01:00:00+00:00",
                "state": "Draft",
                "review_status": "unreviewed",
                "provenance": "manual_unit_sell",
                "currency": "AUD",
                "tax_treatment": "excluded_not_calculated",
                "calculation_version": 1,
                "scope": scope,
                "system_match": None,
                "lines": [],
                "summary": estimates.summary_for(scope, []),
            }
        )
        report.pop("scope")
        report.update(
            schema_version=estimate_reports.COMPLETE_SCHEMA_VERSION
            if profile == "complete"
            else estimate_reports.REPORT_SCHEMA_VERSION,
            render_version=5 if profile == "complete" else 4,
            estimate=estimate,
        )
    return _seal(report)


@pytest.mark.parametrize("profile", ["scope-only", "scope-and-system", "estimate-only", "complete"])
def test_all_profiles_render_entity_history_without_relabeling_it_as_observations(profile):
    snapshot = report_snapshot(profile)
    before = deepcopy(snapshot)
    if "estimate" in snapshot:
        pdf = estimate_outputs.render_estimate_report_pdf(snapshot)
        workbook = estimate_outputs.render_estimate_report_xlsx(snapshot)
    else:
        pdf = scope_outputs.render_scope_report_pdf(snapshot)
        workbook = scope_outputs.render_scope_report_xlsx(snapshot)
    text = " ".join(_pdf_text(pdf).split())
    book = load_workbook(io.BytesIO(workbook), data_only=False)
    cells = [cell for sheet in book for row in sheet for cell in row]
    strings = " ".join(str(cell.value) for cell in cells if cell.value is not None)
    for value in (
        "Defect: D-01",
        "Opening: O-01 Shared wall",
        "Item changed since page review; review again",
        "Item removed since page review; historical page reference retained",
        "Imported source and review claims are unverified",
        "no technical or physical approval",
        "Synthetic <report>.pdf",
    ):
        assert value in text
        assert value in strings
    for identity in (_id(4), _id(5), _id(8), _id(9), _id(10)):
        assert identity in strings
    assert all(cell.data_type != "f" and cell.hyperlink is None for cell in cells)
    assert snapshot == before
    if "estimate" not in snapshot:
        assert "Target ID" in [cell.value for row in book["Page Review References"] for cell in row]
    else:
        assert snapshot["estimate"]["summary"]["priced_subtotal_ex_tax"] == "0.00"


@pytest.mark.parametrize("profile", ["scope-only", "scope-and-system", "estimate-only", "complete"])
def test_v4_reports_require_new_renderer_versions_and_preserve_legacy_readers(profile):
    snapshot = report_snapshot(profile)
    service = estimate_reports if "estimate" in snapshot else scope_reports
    service.validate_report_snapshot(snapshot)
    invalid = deepcopy(snapshot)
    invalid["render_version"] = (
        3 if profile == "complete" else 2 if profile == "scope-and-system" else 1
    )
    _seal(invalid)
    with pytest.raises(DraftScopeError, match="SNAPSHOT_INVALID"):
        service.validate_report_snapshot(invalid)
    legacy = report_snapshot(profile, _snapshot(imported=False)["scope"])
    legacy["render_version"] = invalid["render_version"]
    _seal(legacy)
    service.validate_report_snapshot(legacy)


def test_v3_observation_report_retains_legacy_columns():
    snapshot = _snapshot(imported=False)
    scope = snapshot["scope"]
    observation = scope["content"]["observations"][0]
    scope.update(
        schema_version="CLASSIFIRE-DRAFT-SCOPE-v3",
        provenance="evidence_review",
        import_lineage=[],
        evidence_refs=[
            page_reference()
            | {
                "method": "human_page_review",
                "observation_id": observation["id"],
                "observation_sha256": observation_hash(observation),
            }
        ],
    )
    _seal(scope)
    _seal(snapshot)
    book = load_workbook(io.BytesIO(scope_outputs.render_scope_report_xlsx(snapshot)))
    values = [cell.value for row in book["Page Review References"] for cell in row]
    assert "Observation ID" in values and "Target ID" not in values
    assert "saved revision review status" not in values


def test_package_inventory_retains_all_v4_reference_kinds_and_rejects_omission():
    scope = entity_scope()
    members = {"artifacts/scope.json": packages.encode(scope)}
    manifest = {
        "schema_version": packages.SCHEMA,
        "state": "Draft",
        "authority": "historical_only",
        "project": _snapshot(imported=False)["project"],
        "draft_scope_id": scope["artifact_id"],
        "selection": packages.selection({"scope_revision": scope["revision"]}).model_dump(),
        "notice": packages.NOTICE,
        "capabilities": {
            "scope": "included",
            "system_match": "not_selected",
            "estimate": "not_selected",
            "reports": "not_selected",
        },
        "source_manifest": packages.source_manifest(scope, None, None),
        "members": [
            {"path": name, "sha256": packages.digest(raw), "size_bytes": len(raw)}
            for name, raw in members.items()
        ],
        "package_id": _id(60),
        "revision": 1,
        "parent_hash": None,
        "created_by": _id(1),
        "created_at": "2026-09-06T01:00:00+00:00",
    }
    original = packages._archive(manifest, members)
    inspected = imports.inspect_package(original)
    assert inspected.scope == scope
    assert inspected.members["artifacts/scope.json"] == members["artifacts/scope.json"]
    assert len(inspected.manifest["source_manifest"]) == len(scope["evidence_refs"])
    assert all(ref["membership"] == "external" for ref in inspected.manifest["source_manifest"])
    manifest["source_manifest"].pop()
    with pytest.raises(packages.PackageError, match="PACKAGE_CONTENT_INVALID"):
        imports.inspect_package(packages._archive(manifest, members))


def test_shared_scope_template_displays_entity_status_and_escapes_labels():
    from classifire.ui import templates

    scope = entity_scope()
    rendered = templates.env.get_template("draft_scope_report_content.html").render(envelope=scope)
    assert "Item changed since page review; review again" in rendered
    assert "Item removed since page review; historical page reference retained" in rendered
    assert "Imported source and review claims are unverified" in rendered
    assert "Synthetic &lt;report&gt;.pdf" in rendered
    assert '<img src="https://example.invalid/source.png"/>' not in rendered
