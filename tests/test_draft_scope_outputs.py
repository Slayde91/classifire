from __future__ import annotations

import hashlib
import io
import json
from copy import deepcopy
from decimal import Decimal
from typing import Any
from uuid import UUID

import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pypdf import PdfReader

from classifire.outputs.draft_scope import render_scope_report_pdf, render_scope_report_xlsx


def _id(number: int) -> str:
    return str(UUID(int=number))


def _seal(value: dict[str, Any]) -> dict[str, Any]:
    value.pop("sha256", None)
    value["sha256"] = hashlib.sha256(
        json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode()
    ).hexdigest()
    return value


def _snapshot(*, imported: bool = True, long_text: bool = False) -> dict[str, Any]:
    note = (
        "Long note starts. " + "Evidence remains unresolved. " * 130 + "END-OF-LONG-NOTE"
        if long_text
        else 'Do not run <img src="https://example.invalid/source.png"/> or <b>approve</b>.'
    )
    scope = _seal(
        {
            "schema_version": "CLASSIFIRE-DRAFT-SCOPE-v1",
            "artifact_id": _id(3),
            "project_id": _id(2),
            "revision": 1,
            "parent_hash": None,
            "created_by": _id(1),
            "created_at": "2026-09-05T00:00:00+00:00",
            "state": "Draft",
            "provenance": "manual",
            "review_status": "unreviewed",
            "content": {
                "defects": [{"id": _id(4), "label": "D-01", "description": note}],
                "openings": [
                    {
                        "id": _id(5),
                        "label": "O-01 Shared wall",
                        "defect_id": _id(4),
                        "plane": "wall",
                        "substrate": "Masonry",
                        "width_mm": "250.125",
                        "height_mm": "150",
                        "blank": False,
                        "state": "Inferred",
                    },
                    {
                        "id": _id(6),
                        "label": "O-02 Floor",
                        "defect_id": _id(4),
                        "plane": "floor",
                        "substrate": "Concrete",
                        "width_mm": None,
                        "height_mm": None,
                        "blank": False,
                        "state": "Unresolved",
                    },
                    {
                        "id": _id(7),
                        "label": "O-03 Blank",
                        "defect_id": _id(4),
                        "plane": "unknown",
                        "substrate": "",
                        "width_mm": None,
                        "height_mm": None,
                        "blank": True,
                        "state": "Unresolved",
                    },
                ],
                "services": [
                    {
                        "id": _id(8),
                        "label": '=HYPERLINK("https://example.invalid", "click")',
                        "opening_ids": [_id(5), _id(6)],
                        "service_type": "Cable bundle",
                        "quantity": "2.125",
                        "unit": "m",
                        "state": "Confirmed",
                    },
                    {
                        "id": _id(9),
                        "label": "Pipe",
                        "opening_ids": [_id(5)],
                        "service_type": "https://example.invalid/pipe",
                        "quantity": None,
                        "unit": "each",
                        "state": "Unresolved",
                    },
                ],
                "observations": [{"id": _id(10), "text": note, "state": "Unresolved"}],
                "assumptions": ["No technical compatibility is asserted."],
                "exclusions": ["Technical selection and estimating are unavailable."],
            },
        }
    )
    if imported:
        source = {
            key: scope[key]
            for key in (
                "schema_version",
                "artifact_id",
                "project_id",
                "revision",
                "created_by",
                "created_at",
                "sha256",
            )
        }
        source["file_sha256"] = "a" * 64
        scope.update(
            schema_version="CLASSIFIRE-DRAFT-SCOPE-v2",
            provenance="imported",
            import_lineage=[source],
        )
        _seal(scope)
    return _seal(
        {
            "schema_version": "CLASSIFIRE-DRAFT-SCOPE-REPORT-v1",
            "report_id": _id(20),
            "project": {
                "id": _id(2),
                "reference": "CF-SYN-REPORT",
                "name": "Synthetic scope <draft>",
            },
            "scope": scope,
            "profile": "scope-only",
            "render_version": 1,
            "created_by": _id(1),
            "created_at": "2026-09-05T01:00:00+00:00",
            "state": "Draft",
            "review_status": "unreviewed",
        }
    )


def _pdf_text(raw: bytes) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(raw)).pages)


@pytest.mark.parametrize("imported", [False, True])
def test_both_reports_preserve_identity_scope_and_explicit_unknowns(imported: bool) -> None:
    snapshot = _snapshot(imported=imported)
    original = deepcopy(snapshot)
    pdf = render_scope_report_pdf(snapshot)
    workbook = load_workbook(io.BytesIO(render_scope_report_xlsx(snapshot)))
    text = _pdf_text(pdf)

    assert snapshot == original
    assert snapshot["report_id"] in text
    assert snapshot["sha256"] in text
    assert snapshot["scope"]["sha256"] in text
    assert "Synthetic scope <draft>" in text
    assert "DRAFT - UNREVIEWED" in text
    assert "2.125 m" in text
    assert "Unknown / unavailable" in text
    assert "O-03 Blank" in text
    assert "Confirmed - unreviewed" in text
    assert 'Do not run <img src="https://example.invalid/source.png"/>' in text
    assert "No technical compatibility is asserted." in text
    assert "Technical selection and estimating are unavailable." in text
    for number in range(4, 11):
        assert _id(number) in text.replace("\n", "")
    if imported:
        assert "unverified declarations" in text
        assert "a" * 64 in text

    assert set(workbook.sheetnames) == {
        "Summary",
        "Defects",
        "Openings",
        "Services",
        "Service Links",
        "Observations",
        "Assumptions",
        "Exclusions",
        "Import History",
    }
    assert workbook["Summary"]["B1"].value is None
    assert workbook["Summary"]["C1"].value == "CLASSIFIRE"
    values = [cell.value for sheet in workbook for row in sheet for cell in row]
    assert snapshot["sha256"] in values
    assert snapshot["scope"]["sha256"] in values
    for sheet in workbook:
        assert sheet.freeze_panes is not None
        assert sheet.auto_filter.ref is not None
    links = list(workbook["Service Links"].iter_rows(min_row=7, values_only=True))
    assert [row[1:3] for row in links] == [(_id(8), _id(5)), (_id(8), _id(6)), (_id(9), _id(5))]
    assert workbook["Services"]["E7"].data_type == "n"
    assert Decimal(str(workbook["Services"]["E7"].value)) == Decimal("2.125")
    assert workbook["Services"]["E8"].value == "Unknown / unavailable"
    assert workbook["Openings"]["G7"].data_type == "n"
    assert Decimal(str(workbook["Openings"]["G7"].value)) == Decimal("250.125")
    assert workbook["Openings"]["I9"].value is True


def test_spreadsheet_does_not_execute_or_link_submitted_text() -> None:
    snapshot = _snapshot()
    workbook = load_workbook(io.BytesIO(render_scope_report_xlsx(snapshot)))
    assert workbook["Services"]["C7"].value == snapshot["scope"]["content"]["services"][0]["label"]
    assert workbook["Services"]["D8"].value == "https://example.invalid/pipe"
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                assert cell.data_type != "f"
                assert cell.hyperlink is None


def test_long_text_flows_across_pdf_pages_and_readable_workbook_continuations() -> None:
    snapshot = _snapshot(long_text=True)
    pdf = render_scope_report_pdf(snapshot)
    text = _pdf_text(pdf)
    assert len(PdfReader(io.BytesIO(pdf)).pages) >= 3
    assert text.count("END-OF-LONG-NOTE") == 2
    # Extractors insert line/page breaks between words; retain every repeated note.
    assert text.count("remains") == 260
    assert text.count("unresolved.") == 260
    workbook = load_workbook(io.BytesIO(render_scope_report_xlsx(snapshot)))
    parts = [row[3] for row in workbook["Defects"].iter_rows(min_row=7, values_only=True)]
    assert (
        "".join(value for value in parts if isinstance(value, str))
        == snapshot["scope"]["content"]["defects"][0]["description"]
    )
    assert len(parts) > 1
    assert all(
        dimension.height <= 150
        for dimension in workbook["Defects"].row_dimensions.values()
        if dimension.height
    )


def test_unicode_and_control_characters_are_not_silently_lost() -> None:
    snapshot = _snapshot()
    original = "Caf\u00e9 \u03b1 \u0416 \u5899\u5b57 \U0001f642 unsafe\x01control"
    snapshot["scope"]["content"]["observations"][0]["text"] = original
    _seal(snapshot["scope"])
    _seal(snapshot)
    text = _pdf_text(render_scope_report_pdf(snapshot))
    assert "Caf\u00e9 \u03b1 \u0416 \u5899\u5b57" in text
    assert "[U+1F642]" in text
    assert "unsafe[U+0001]control" in text
    workbook = load_workbook(io.BytesIO(render_scope_report_xlsx(snapshot)))
    assert workbook["Observations"]["C7"].value == original.replace("\x01", "[U+0001]")


@pytest.mark.parametrize("renderer", [render_scope_report_pdf, render_scope_report_xlsx])
def test_renderers_refuse_tampered_snapshot_and_source(renderer: Any) -> None:
    snapshot = _snapshot()
    snapshot["project"]["name"] = "Changed after capture"
    with pytest.raises(ValueError):
        renderer(snapshot)
    snapshot = _snapshot()
    snapshot["scope"]["content"]["services"][0]["quantity"] = "99"
    _seal(snapshot)
    with pytest.raises(ValueError):
        renderer(snapshot)
