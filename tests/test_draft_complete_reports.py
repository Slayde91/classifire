from __future__ import annotations

import copy
import io
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from test_draft_constraint_review import base_case as _base_case
from test_draft_constraint_review import case as _case
from test_draft_constraint_review import inputs
from test_draft_estimate_reports import priced
from test_draft_estimate_ui import _disable_downstream
from test_draft_estimates import add_payload, prepare
from test_draft_scope import setup, uid  # noqa: F401
from test_draft_service_size_review import size_inputs
from test_draft_system_matches import counts, create

from classifire.models import User
from classifire.outputs import draft_estimate as outputs
from classifire.services import draft_estimate_reports as reports
from classifire.services import draft_estimates as estimates
from classifire.services.draft_scope import DraftScopeError
from classifire.services.draft_system_matches import save_constraint_review, save_review

base_case = _base_case
case = _case


def pdf_text(content):
    return "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(content)).pages)


@pytest.mark.parametrize("version", [1, 2, 3])
def test_complete_pair_retains_scope_technical_commercial_and_history(case, version, monkeypatch):
    with case["factory"]() as db:
        actor = db.get(User, uid(100))
        match = create(db, case)
        if version > 1:
            save_constraint_review(
                db,
                actor,
                case["draft_id"],
                match.id,
                1,
                case["variant_id"],
                size_inputs() if version == 3 else inputs(),
                storage_root=case["storage_root"],
                service_size=version == 3,
            )
        estimate = estimates.create_estimate(
            db,
            actor,
            case["draft_id"],
            2,
            match_id=match.id,
            match_revision=1 if version == 1 else 2,
        )
        envelope = estimates.add_line(
            db,
            actor,
            case["draft_id"],
            estimate.id,
            1,
            add_payload(source_note='=HYPERLINK("bad") <script> literal evidence'),
        )
        old = reports.create_report(db, actor, case["draft_id"], estimate.id, 2)
        old_pdf = reports.report_bytes(db, actor, case["draft_id"], estimate.id, old.id, "pdf")
        before = counts(db)
        _disable_downstream(monkeypatch)
        row = reports.create_report(db, actor, case["draft_id"], estimate.id, 2, profile="complete")
        ids = (case["draft_id"], estimate.id, row.id)
        snapshot = reports.read_report(db, actor, *ids)
        assert snapshot["schema_version"] == reports.COMPLETE_SCHEMA_VERSION
        assert snapshot["profile"] == "complete" and snapshot["render_version"] == 3
        assert snapshot["estimate"] == envelope
        pdf = reports.report_bytes(db, actor, *ids, "pdf")
        xlsx = reports.report_bytes(db, actor, *ids, "xlsx")
        text = pdf_text(pdf)
        for value in (
            "Complete Draft Report",
            "Scope overview",
            "Saved technical review summary",
            "Retained technical evidence and decisions",
            "2.01",
            "1.005",
            "Unapproved candidate review",
            "Site verification required",
        ):
            assert value in text
        if version == 3:
            assert "Measured Service Outside Diameter Range" in text
        book = load_workbook(io.BytesIO(xlsx), data_only=False)
        assert {
            "Scope context",
            "Lines",
            "History",
            "Technical summary",
            "System candidates",
            "Measured limits",
            "Review and coverage",
        } <= set(book.sheetnames)
        assert book["Lines"]["G7"].value == "2.01"
        assert book["Lines"]["H7"].value == 2.01
        assert book["Lines"]["G7"].data_type == "s"
        for sheet in book:
            assert sheet.auto_filter.ref and sheet.freeze_panes == "C7"
            assert not any(c.data_type == "f" or c.hyperlink for r in sheet for c in r)
        with ZipFile(io.BytesIO(xlsx)) as z:
            assert z.read("xl/media/image1.png") == outputs._logo_path().read_bytes()
        book.close()
        for table in (
            "estimates",
            "openings",
            "services",
            "physical_model_locks",
            "draft_scope_revisions",
            "draft_system_match_revisions",
        ):
            assert counts(db)[table] == before[table]
        estimates.override_line(
            db,
            actor,
            case["draft_id"],
            estimate.id,
            2,
            envelope["lines"][0]["line_id"],
            {"quantity": "3", "unit_sell_rate": "2", "reason": "Later input"},
        )
        saved_match = envelope["system_match"]
        save_review(
            db, actor, case["draft_id"], match.id, saved_match["revision"], saved_match["decisions"]
        )
        db.commit()
        stale = reports.report_staleness(db, actor, *ids, storage_root=case["storage_root"])
        assert {"REPORT_ESTIMATE_CHANGED", "ESTIMATE_REVIEW_CHANGED"} <= set(stale)
        case["source_path"].write_bytes(b"Synthetic later source")
        assert reports.report_staleness(db, actor, *ids, storage_root=case["storage_root"])
        assert (
            reports.report_bytes(db, actor, case["draft_id"], estimate.id, old.id, "pdf") == old_pdf
        )
        db.commit()
    with case["factory"]() as db:
        actor = db.get(User, uid(100))
        assert reports.report_bytes(db, actor, *ids, "pdf") == pdf
        assert reports.report_bytes(db, actor, *ids, "xlsx") == xlsx
        actor.role = "project_manager"
        db.commit()
        for read in (
            lambda: reports.read_report(db, actor, *ids),
            lambda: reports.report_bytes(db, actor, *ids, "xlsx"),
        ):
            with pytest.raises(DraftScopeError) as error:
                read()
            assert error.value.status_code == 403


def test_empty_complete_profile_explicitly_reports_missing_sections(setup):  # noqa: F811
    with setup() as db:
        actor, draft, estimate = prepare(db)
        row = reports.create_report(db, actor, draft.id, estimate.id, 1, profile="complete")
        text = pdf_text(reports.report_bytes(db, actor, draft.id, estimate.id, row.id, "pdf"))
        assert "Unavailable - no System Match attached" in text
        assert "Unavailable - no Estimate work lines" in text
        assert "zero subtotal is not a quote" in text
        book = load_workbook(
            io.BytesIO(reports.report_bytes(db, actor, draft.id, estimate.id, row.id, "xlsx"))
        )
        assert "Unavailable - no System Match attached" in book["Technical summary"]["C7"].value
        assert "System candidates" not in book.sheetnames
        book.close()


@pytest.mark.parametrize(
    "change",
    [
        {"profile": "canonical"},
        {"profile": "estimate-only"},
        {"schema_version": reports.REPORT_SCHEMA_VERSION},
        {"render_version": 2},
        {"review_status": "approved"},
    ],
)
def test_complete_contract_cannot_reseal_mixed_version_or_authority(setup, change):  # noqa: F811
    with setup() as db:
        actor, draft, estimate, _ = priced(db)
        row = reports.create_report(db, actor, draft.id, estimate.id, 2, profile="complete")
        snapshot = copy.deepcopy(reports.read_report(db, actor, draft.id, estimate.id, row.id))
        snapshot.update(change)
        snapshot["sha256"] = reports._checksum(snapshot)
        with pytest.raises(DraftScopeError):
            reports.validate_report_snapshot(snapshot)
        with pytest.raises(DraftScopeError):
            reports.create_report(db, actor, draft.id, estimate.id, 2, profile="canonical")
