from __future__ import annotations

import io
import json
from copy import deepcopy

import pytest
from openpyxl import load_workbook
from test_draft_constraint_review import base_case as _base_case
from test_draft_constraint_review import case as _case
from test_draft_constraint_review import inputs
from test_draft_entity_evidence_outputs import report_snapshot
from test_draft_estimates import add_payload
from test_draft_multi_review_report_outputs import collection_snapshot
from test_draft_scope import uid
from test_draft_scope_outputs import _pdf_text, _seal, _snapshot
from test_draft_system_matches import create
from test_draft_word_evidence_outputs import word_scope

from classifire.models import User
from classifire.outputs import draft_estimate as outputs
from classifire.services import draft_estimate_reports as reports
from classifire.services import draft_estimates as estimates
from classifire.services import draft_scope_reports
from classifire.services import draft_system_matches as matches

base_case = _base_case
case = _case


def complete_snapshot(*, bound=True):
    scoped = collection_snapshot(word=True)
    snapshot = report_snapshot("complete", scoped["scope"])
    if bound:
        snapshot["estimate"]["system_match"] = scoped["system_matches"][1]
        _seal(snapshot["estimate"])
    snapshot.update(
        schema_version=reports.MULTI_COMPLETE_SCHEMA_VERSION,
        render_version=12,
        system_matches=scoped["system_matches"],
    )
    return _seal(snapshot)


def cells(book):
    return [cell for sheet in book for row in sheet for cell in row if cell.value is not None]


def assert_context(snapshot, pdf, book):
    text = " ".join(_pdf_text(pdf).split())
    values = " ".join(str(cell.value) for cell in cells(book))
    rows = {}
    for _, key, label, value in book["Technical summary"].iter_rows(min_row=7, values_only=True):
        rows.setdefault(key, {})[label] = value
    assert len(rows) == len(snapshot["system_matches"])
    bound = snapshot["estimate"]["system_match"]
    for review in snapshot["system_matches"]:
        row = rows[f"{review['artifact_id']} / revision {review['revision']}"]
        assert row["Review SHA-256"] == review["sha256"]
        assert row["Opening ID"] == review["target"]["opening_id"]
        assert row["Service ID"] == (review["target"]["service_id"] or "None - blank opening")
        assert ("Estimate-bound" in row["Relationship to Estimate"]) is (review == bound)
        for identity in (review["artifact_id"], review["sha256"], review["target"]["opening_id"]):
            assert identity in text and identity in values
    for expected in (
        "Additional report-only context - not a pricing input",
        "Review records are not work quantities",
        "Blank opening - no Service",
        "The Estimate and its prices are unchanged",
        "Unapproved",
        "Tax not calculated",
        snapshot["estimate"]["sha256"],
        snapshot["sha256"],
    ):
        assert expected in text or (
            expected == "Tax not calculated" and "Tax: Not calculated" in text
        )
        assert expected in values
    assert all(cell.data_type != "f" and cell.hyperlink is None for cell in cells(book))
    for sheet in book:
        assert sheet.auto_filter.ref and sheet.freeze_panes == "C7"
    return text, values


@pytest.mark.parametrize("bound", [True, False], ids=["bound-review", "report-context-only"])
def test_complete_word_pair_preserves_exact_estimate_and_all_selected_row_context(bound, tmp_path):
    snapshot = complete_snapshot(bound=bound)
    reports.validate_report_snapshot(snapshot)
    original = deepcopy(snapshot)
    estimate_bytes = reports._canonical(snapshot["estimate"])
    pdf = outputs.render_estimate_report_pdf(snapshot)
    xlsx = outputs.render_estimate_report_xlsx(snapshot)
    assert snapshot == original
    assert reports._canonical(snapshot["estimate"]) == estimate_bytes
    book = load_workbook(io.BytesIO(xlsx), data_only=False)
    text, values = assert_context(snapshot, pdf, book)
    for expected in (
        "Synthetic Word.docx",
        "body-2/row-1/cell-1/p-1",
        "word/media/image1.png",
        "untrusted Word text",
        "synthetic-page-suggester",
        "Unknown / unavailable",
        'Do not run <img src="https://example.invalid/source.png"/>',
    ):
        assert expected in text and expected in values
    assert snapshot["estimate"]["summary"]["priced_subtotal_ex_tax"] == "0.00"
    assert not snapshot["estimate"]["lines"]
    assert "zero subtotal is not a quote" in text and "zero subtotal is not a quote" in values
    selected = snapshot["system_matches"]
    assert selected[0]["target"]["service_id"] == selected[1]["target"]["service_id"]
    assert selected[0]["target"]["opening_id"] != selected[1]["target"]["opening_id"]
    assert selected[2]["target"]["blank_opening"]
    assert len(pdf) < draft_scope_reports.MAX_OUTPUT_BYTES
    assert len(xlsx) < draft_scope_reports.MAX_OUTPUT_BYTES
    book.close()
    (tmp_path / "complete-word-review.pdf").write_bytes(pdf)
    (tmp_path / "complete-word-review.xlsx").write_bytes(xlsx)
    (tmp_path / "complete-word-snapshot.json").write_text(json.dumps(snapshot, ensure_ascii=False))


def test_complete_context_does_not_change_saved_amounts_overrides_or_unknown_work(
    case, monkeypatch, tmp_path
):
    with case["factory"]() as db:
        actor = db.get(User, uid(100))
        selected = []
        for opening, service in ((uid(2), uid(5)), (uid(3), uid(5)), (uid(4), None)):
            row = create(db, case, opening_id=opening, service_id=service)
            review = matches.read_match_revision(db, actor, case["draft_id"], row.id, 1)
            if not selected:
                review = matches.save_constraint_review(
                    db,
                    actor,
                    case["draft_id"],
                    row.id,
                    1,
                    case["variant_id"],
                    inputs(None, None, None),
                    storage_root=case["storage_root"],
                )
            selected.append(review)
        bound = selected[0]
        estimate = estimates.create_estimate(
            db,
            actor,
            case["draft_id"],
            2,
            match_id=bound["artifact_id"],
            match_revision=bound["revision"],
        )
        saved = estimates.add_line(db, actor, case["draft_id"], estimate.id, 1, add_payload())
        saved = estimates.override_line(
            db,
            actor,
            case["draft_id"],
            estimate.id,
            2,
            saved["lines"][0]["line_id"],
            dict(
                quantity="3",
                unit_sell_rate="2.345678",
                reason='Manual override <b>literal</b> =HYPERLINK("bad")',
            ),
        )
        saved = estimates.add_line(
            db,
            actor,
            case["draft_id"],
            estimate.id,
            3,
            add_payload(target_id=uid(6), quantity=None, unit_sell_rate=None),
        )
        saved = estimates.add_line(
            db,
            actor,
            case["draft_id"],
            estimate.id,
            4,
            add_payload(
                target_kind="blank_opening",
                target_id=uid(4),
                unit="m2",
                quantity="1.5",
                unit_sell_rate="4.123456",
                reason="Synthetic measured blank area",
            ),
        )
        assert saved["summary"]["priced_subtotal_ex_tax"] == "13.23"
        assert len(saved["summary"]["unpriced_line_ids"]) == 1
        original_bytes = reports._canonical(saved)
        legacy = reports.create_report(
            db, actor, case["draft_id"], estimate.id, 5, profile="complete"
        )
        legacy_book = load_workbook(io.BytesIO(legacy.xlsx_bytes))

        def forbidden(*args, **kwargs):
            pytest.fail("Report context must not rerun an upstream capability")

        for target in (
            "classifire.services.draft_system_matches.create_match",
            "classifire.services.technical.search_variants",
            "classifire.services.calculation.recalculate_estimate",
            "classifire.services.snapshot.lock_snapshot",
            "classifire.services.draft_estimates.create_estimate",
        ):
            monkeypatch.setattr(target, forbidden)
        report = reports.create_report(
            db,
            actor,
            case["draft_id"],
            estimate.id,
            5,
            profile="complete",
            matches=[
                dict(match_id=item["artifact_id"], match_revision=item["revision"])
                for item in reversed(selected)
            ],
        )
        snapshot = reports.read_report(db, actor, case["draft_id"], estimate.id, report.id)
        assert reports._canonical(snapshot["estimate"]) == original_bytes
        assert snapshot["estimate"] == saved
        assert snapshot["system_matches"] == sorted(selected, key=lambda item: item["artifact_id"])
        pdf, xlsx = report.pdf_bytes, report.xlsx_bytes
        book = load_workbook(io.BytesIO(xlsx))
        text, values = assert_context(snapshot, pdf, book)
        for expected in (
            "13.23",
            "7.04",
            "6.19",
            "2.345678",
            "4.123456",
            "Unknown / unavailable",
            'Manual override <b>literal</b> =HYPERLINK("bad")',
            "unresolved",
        ):
            assert expected in text and expected in values
        for name in ("Lines", "Line details", "History", "Coverage", "Scope context"):
            current_rows, prior_rows = list(book[name].values), list(legacy_book[name].values)
            assert current_rows[2][1] == f"Estimate revision 5 | Report {report.id}"
            assert prior_rows[2][1] == f"Estimate revision 5 | Report {legacy.id}"
            # Distinct report records have distinct identity headers; every other cell is equal.
            assert current_rows[:2] + current_rows[3:] == prior_rows[:2] + prior_rows[3:]
        line_rows = {row[1]: row for row in book["Lines"].iter_rows(min_row=7, values_only=True)}
        for line in saved["lines"]:
            row = line_rows[line["line_id"]]
            assert row[4:7] == tuple(
                line[key] if line[key] is not None else "Unknown / unavailable"
                for key in ("quantity", "unit_sell_rate", "subtotal_ex_tax")
            )
        shared = next(line for line in saved["lines"] if line["target_id"] == uid(5))
        assert shared["opening_ids"] == [uid(2), uid(3)]
        assert shared["original_quantity"] == "2" and shared["original_rate"] == "1.005"
        assert snapshot["estimate"]["system_match"] == bound
        assert "must-never-export-source-json" not in text + values
        legacy_book.close()
        book.close()
        db.commit()
    with case["factory"]() as db:
        actor = db.get(User, uid(100))
        assert (
            reports.report_bytes(db, actor, case["draft_id"], estimate.id, report.id, "pdf") == pdf
        )
        assert (
            reports.report_bytes(db, actor, case["draft_id"], estimate.id, report.id, "xlsx")
            == xlsx
        )
    (tmp_path / "complete-priced-review.pdf").write_bytes(pdf)
    (tmp_path / "complete-priced-review.xlsx").write_bytes(xlsx)
    (tmp_path / "complete-priced-snapshot.json").write_text(
        json.dumps(snapshot, ensure_ascii=False)
    )


@pytest.mark.parametrize(
    "word,complete", [(False, False), (False, True), (True, False), (True, True)]
)
def test_legacy_estimate_and_complete_layout_remains_unchanged(word, complete):
    snapshot = report_snapshot(
        "complete" if complete else "estimate-only",
        word_scope() if word else _snapshot(imported=False)["scope"],
    )
    snapshot["render_version"] = (11 if complete else 10) if word else (3 if complete else 1)
    _seal(snapshot)
    reports.validate_report_snapshot(snapshot)
    assert outputs._complete_reviews(snapshot) == []
    pdf = outputs.render_estimate_report_pdf(snapshot)
    book = load_workbook(io.BytesIO(outputs.render_estimate_report_xlsx(snapshot)))
    assert "Additional report-only context" not in _pdf_text(pdf)
    if complete:
        assert [cell.value for cell in book["Technical summary"][6]] == [
            "Part",
            "Field",
            "Saved finding",
        ]
    else:
        assert "Technical summary" not in book.sheetnames
    book.close()
